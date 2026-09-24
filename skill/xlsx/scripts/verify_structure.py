#!/usr/bin/env python3
"""Verify that editing a workbook changed only the cells the edit was meant to change.

Compares the user's original file with the edited output and fails when anything outside the
declared edit ranges differs: sheets, merged ranges, column widths, row heights, hidden rows and
columns, page setup, margins, page header and footer, print titles, page breaks, freeze panes,
data validation, conditional formatting, images and charts, or any cell's value or formatting.

It also fails when a cell's content no longer fits the widths and heights the sheet already has:
a number that prints as ####, text cut off at the cell edge, wrapped text taller than its row.
A cell that already overflowed in the original is not reported. Run it after recalc.py, so
formula results are measured too.

Usage:
    python verify_structure.py ORIGINAL OUTPUT --allow 'Sheet 2!B12:D14' --allow 'Sheet 2!A19'

Ranges accept a cell (A1), a block (A1:D9), whole rows (12:14) or whole columns (B:D). Quote the
sheet name or not: 'Sheet 2'!A1 and Sheet 2!A1 are both accepted.

A tab the edit adds must be declared with --new-sheet. 'NAME:TEMPLATE' declares a tab copied from
an existing sheet: it is held to that template exactly like an edit, so everything outside its
--allow ranges must still match the template. A bare NAME is only permitted to exist.

    python verify_structure.py ORIGINAL OUTPUT --new-sheet 'Investigation:Sheet1' \
      --allow 'Investigation!B8:F12' --allow 'Investigation!C22'

Prints JSON and exits 0 on success, 2 when unexpected changes are found, 1 on bad input.
"""
import argparse
import json
import re
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
from openpyxl.utils.exceptions import CellCoordinatesException

from fit import base_font_name, base_font_size, fit_problems
from layout import layout_key

DEFAULT_COLUMN_WIDTH = 8.43
DEFAULT_ROW_HEIGHT = 15.0
COLUMN_WIDTH_TOLERANCE = 0.75
# LibreOffice stores row heights in 0.75pt steps, so a re-save shifts them by up to that much.
ROW_HEIGHT_TOLERANCE = 1.0
RANGE_OPERATORS = {'between', 'notBetween'}
AUTOMATIC_COLOR_INDEX = 64
MAX_ISSUES = 100
UNBOUNDED = 1_048_576


class InputError(Exception):
    pass


def parse_allow(spec):
    """Turn "Sheet!RANGE" into (sheet, min_col, min_row, max_col, max_row)."""
    if '!' not in spec:
        raise InputError(f"--allow must name the sheet, e.g. 'Sheet 2!B12:D14': {spec}")
    sheet, ref = spec.rsplit('!', 1)
    sheet = sheet.strip().strip("'")
    ref = ref.strip().replace('$', '').upper()
    if not sheet or not ref:
        raise InputError(f'Invalid --allow range: {spec}')
    start, _, end = ref.partition(':')
    end = end or start
    try:
        if start.isdigit() and end.isdigit():
            return sheet, 1, int(start), UNBOUNDED, int(end)
        if start.isalpha() and end.isalpha():
            return sheet, column_index_from_string(start), 1, column_index_from_string(end), UNBOUNDED
        start_col, start_row = coordinate_from_string(start)
        end_col, end_row = coordinate_from_string(end)
    except (ValueError, CellCoordinatesException) as error:
        raise InputError(f'Invalid --allow range {spec}: {error}') from error
    min_col, max_col = sorted((column_index_from_string(start_col), column_index_from_string(end_col)))
    min_row, max_row = sorted((start_row, end_row))
    return sheet, min_col, min_row, max_col, max_row


def parse_new_sheet(spec):
    """Turn "NAME" or "NAME:TEMPLATE" into (name, template or None).

    Excel forbids ':' in sheet names, so it cannot be part of either name.
    """
    name, _, template = spec.partition(':')
    name, template = name.strip().strip("'"), template.strip().strip("'")
    if not name or (':' in spec and not template):
        raise InputError(f"--new-sheet takes 'NAME' or 'NAME:TEMPLATE': {spec}")
    return name, template or None


def is_allowed(allowed, sheet, row, column):
    return any(
        name == sheet and min_col <= column <= max_col and min_row <= row <= max_row
        for name, min_col, min_row, max_col, max_row in allowed
    )


def color_key(color):
    """Compare colors by meaning: LibreOffice rewrites an ARGB alpha byte, and saves Excel's
    automatic color (indexed 64, the system foreground) as `auto` or as no color at all."""
    if color is None:
        return None
    if color.type == 'rgb' and isinstance(color.rgb, str):
        return color.rgb[-6:].upper()
    if color.type == 'theme':
        return f'theme:{color.theme}:{round(color.tint or 0, 4)}'
    if color.type == 'indexed' and color.indexed != AUTOMATIC_COLOR_INDEX:
        return f'indexed:{color.indexed}'
    return None


def style_key(cell):
    font, fill, border, alignment = cell.font, cell.fill, cell.border, cell.alignment
    fill_key = None
    if fill is not None and fill.fill_type:
        fill_key = (fill.fill_type, color_key(fill.fgColor))
    return (
        (font.name, float(font.sz or 0), bool(font.b), bool(font.i), font.u, bool(font.strike), color_key(font.color)),
        fill_key,
        tuple(
            (side.style, color_key(side.color) if side.style else None)
            for side in (border.left, border.right, border.top, border.bottom)
        ),
        # An unset alignment and Excel's defaults render the same; LibreOffice writes the defaults out.
        (alignment.horizontal or 'general', alignment.vertical or 'bottom', bool(alignment.wrap_text), alignment.indent or 0),
        cell.number_format,
    )


def value_key(value):
    if value is None or value == '':
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return round(float(value), 9)
    return value


def media_counts(path):
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
    return {
        'images': sum(1 for name in names if name.startswith('xl/media/')),
        'charts': sum(1 for name in names if re.match(r'xl/charts/chart\d+\.xml$', name)),
        'drawings': sum(1 for name in names if re.match(r'xl/drawings/drawing\d+\.xml$', name)),
    }


def dimension_changes(kind, original, output, default, tolerance):
    """Differences that exceed tolerance; an unset size counts as the sheet default."""
    changes = []
    for key in sorted(set(original) | set(output), key=str):
        before = original.get(key) or default
        after = output.get(key) or default
        if abs(before - after) > tolerance:
            changes.append(f'{kind} {key}: {before} -> {after}')
    return changes


def compare_sheet(name, before, after, allowed, report):
    """Compare `after` (reported as `name`) against `before`, skipping cells `allowed` covers."""
    before_merges = {str(r) for r in before.merged_cells.ranges}
    after_merges = {str(r) for r in after.merged_cells.ranges}
    for ref in sorted(before_merges - after_merges):
        report('unmerged', name, f'merged range {ref} was removed')
    for ref in sorted(after_merges - before_merges):
        report('merged', name, f'merged range {ref} was added')

    widths = dimension_changes(
        'column',
        {k: v.width for k, v in before.column_dimensions.items() if v.width},
        {k: v.width for k, v in after.column_dimensions.items() if v.width},
        DEFAULT_COLUMN_WIDTH,
        COLUMN_WIDTH_TOLERANCE,
    )
    heights = dimension_changes(
        'row',
        {k: v.height for k, v in before.row_dimensions.items() if v.height},
        {k: v.height for k, v in after.row_dimensions.items() if v.height},
        before.sheet_format.defaultRowHeight or DEFAULT_ROW_HEIGHT,
        ROW_HEIGHT_TOLERANCE,
    )
    for detail in widths + heights:
        report('dimension', name, detail)

    before_layout, after_layout = layout_key(before), layout_key(after)
    for setting, value in before_layout.items():
        if value is not None and after_layout[setting] != value:
            report('layout', name, f'{setting}: {value!r} -> {after_layout[setting]!r}')

    def validations(ws):
        # Only "between" rules read a second bound; LibreOffice fills in an unused one on save.
        return {
            (dv.type, dv.operator, str(dv.sqref), dv.formula1, dv.formula2 if dv.operator in RANGE_OPERATORS else None)
            for dv in ws.data_validations.dataValidation
        }

    if validations(before) != validations(after):
        report('data_validation', name, 'data validation rules changed')

    def conditional(ws):
        return {(str(rng.sqref), len(rng.rules)) for rng in ws.conditional_formatting}

    if conditional(before) != conditional(after):
        report('conditional_formatting', name, 'conditional formatting changed')

    coordinates = {(c.row, c.column) for c in before._cells.values()} | {
        (c.row, c.column) for c in after._cells.values()
    }
    checked = changed_in_allowed = 0
    for row, column in sorted(coordinates):
        old, new = before.cell(row=row, column=column), after.cell(row=row, column=column)
        value_changed = value_key(old.value) != value_key(new.value)
        style_changed = style_key(old) != style_key(new)
        if is_allowed(allowed, name, row, column):
            changed_in_allowed += int(value_changed or style_changed)
            continue
        checked += 1
        if value_changed:
            report('value', name, f'{old.coordinate}: {old.value!r} -> {new.value!r}')
        elif style_changed:
            report('style', name, f'{old.coordinate}: formatting changed')
    return checked, changed_in_allowed, len(before_merges)


def check_fit(pairs, before_wb, after_wb, before_values, after_values, report):
    """Report content that outgrew its cell, unless the same cell overflowed in the original."""
    before_size, after_size = base_font_size(before_wb), base_font_size(after_wb)
    before_font, after_font = base_font_name(before_wb), base_font_name(after_wb)
    for name, source in pairs:
        if name not in after_wb.sheetnames:
            continue
        known = fit_problems(before_wb[source], before_values[source], before_size, before_font) if source else {}
        found = fit_problems(after_wb[name], after_values[name], after_size, after_font)
        for coordinate, problem in found.items():
            if coordinate not in known:
                report('fit', name, f'{coordinate}: {problem}')


def check_sheet_list(before_names, after_names, new_sheets, report):
    """Original sheets keep their order; the only extra sheets are the declared new ones."""
    kept = [name for name in after_names if name not in new_sheets]
    if kept != before_names:
        report('sheets', None, f'sheets {before_names} -> {after_names}')
    for name in new_sheets:
        if name not in after_names:
            report('sheets', None, f'declared new sheet {name!r} is not in the output')


def verify(original, output, allowed, new_sheets=None):
    """`new_sheets` maps each tab the edit adds to the sheet it was copied from, or None."""
    new_sheets = new_sheets or {}
    issues = []

    def report(kind, sheet, detail):
        issues.append({'kind': kind, 'sheet': sheet, 'detail': detail})

    before_wb, before_values = load_workbook(original), load_workbook(original, data_only=True)
    after_wb, after_values = load_workbook(output), load_workbook(output, data_only=True)
    clashing = sorted(set(new_sheets) & set(before_wb.sheetnames))
    if clashing:
        raise InputError(f'--new-sheet names sheets that already exist in the original: {clashing}')
    unknown_templates = sorted({t for t in new_sheets.values() if t} - set(before_wb.sheetnames))
    if unknown_templates:
        raise InputError(f'--new-sheet templates that are not in the original: {unknown_templates}')
    unknown = {sheet for sheet, *_ in allowed} - set(before_wb.sheetnames) - set(new_sheets)
    if unknown:
        raise InputError(f'--allow names sheets that are not in the original: {sorted(unknown)}')

    check_sheet_list(before_wb.sheetnames, after_wb.sheetnames, new_sheets, report)

    before_media, after_media = media_counts(original), media_counts(output)
    for kind, count in before_media.items():
        if after_media[kind] < count:
            report('media', None, f'{kind}: {count} in the original, {after_media[kind]} in the output')

    pairs = [(name, name) for name in before_wb.sheetnames] + [
        (name, template) for name, template in new_sheets.items() if template
    ]
    cells_checked = allowed_changes = merges_checked = 0
    for name, source in pairs:
        if name not in after_wb.sheetnames:
            continue
        checked, changed, merges = compare_sheet(name, before_wb[source], after_wb[name], allowed, report)
        cells_checked += checked
        allowed_changes += changed
        merges_checked += merges

    fit_pairs = pairs + [(name, None) for name, template in new_sheets.items() if not template]
    check_fit(fit_pairs, before_wb, after_wb, before_values, after_values, report)

    return {
        'status': 'changes_found' if issues else 'success',
        'sheets_checked': len(pairs),
        'new_sheets': {name: template for name, template in new_sheets.items()},
        'merged_ranges_checked': merges_checked,
        'cells_checked_outside_allowed': cells_checked,
        'cells_changed_inside_allowed': allowed_changes,
        'total_issues': len(issues),
        'issues': issues[:MAX_ISSUES],
        'issues_truncated': max(0, len(issues) - MAX_ISSUES),
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('original', type=Path, help="the user's workbook, before any edit")
    parser.add_argument('output', type=Path, help='the edited workbook you intend to deliver')
    parser.add_argument(
        '--allow',
        action='append',
        default=[],
        metavar="'Sheet!RANGE'",
        help='a range the edit was meant to change; repeat for each range',
    )
    parser.add_argument(
        '--new-sheet',
        action='append',
        default=[],
        metavar="'NAME[:TEMPLATE]'",
        help='a tab the edit adds; with :TEMPLATE it must match that sheet outside its --allow ranges',
    )
    args = parser.parse_args()

    try:
        for path in (args.original, args.output):
            if not path.is_file():
                raise InputError(f'File not found: {path}')
        if args.original.resolve() == args.output.resolve():
            raise InputError('ORIGINAL and OUTPUT are the same file; save the edit under a new name')
        allowed = [parse_allow(spec) for spec in args.allow]
        new_sheets = dict(parse_new_sheet(spec) for spec in args.new_sheet)
        result = verify(args.original, args.output, allowed, new_sheets)
    except InputError as error:
        print(json.dumps({'error': str(error)}, indent=2))
        return 1

    print(json.dumps(result, indent=2))
    return 0 if result['status'] == 'success' else 2


if __name__ == '__main__':
    sys.exit(main())
