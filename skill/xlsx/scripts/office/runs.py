"""The text runs of a DOCX, PPTX or XLSX file, with the font and size each one really gets.

A run's font is rarely on the run itself. Word takes the complex-script font (the one it uses for
Bangla) from the run, then its character style, its paragraph style, the document defaults and
finally the theme; PowerPoint from the run, the text box's list style and the theme. The checks
and fixes for Bangla need the font that is actually used, so this module follows those chains.

`Package` reads an OOXML file into memory and writes it back with only the changed parts
re-serialised, so every other part stays byte for byte as it was.
"""
import re
import zipfile
from pathlib import Path

from lxml import etree

from office import bijoy
from office.bangla import has_bangla

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PKG_REL = 'http://schemas.openxmlformats.org/package/2006/relationships'
CT = 'http://schemas.openxmlformats.org/package/2006/content-types'
NS = {'w': W, 'a': A, 'p': P, 'r': R}
PARSER = etree.XMLParser(resolve_entities=False, no_network=True, remove_blank_text=False)
DOCX_TEXT_PARTS = re.compile(r'^word/(document|header\d*|footer\d*|footnotes|endnotes|comments)\.xml$')
PPTX_TEXT_PARTS = re.compile(r'^ppt/(slides/slide|notesSlides/notesSlide)\d+\.xml$')
MAX_STYLE_DEPTH = 20
DEFAULT_HALF_POINTS = 20
"""Word's size when nothing sets one: 10 pt."""


def w(name):
    return f'{{{W}}}{name}'


def a(name):
    return f'{{{A}}}{name}'


def number_key(name):
    """slide10.xml after slide9.xml."""
    return [int(part) if part.isdigit() else part for part in re.split(r'(\d+)', name)]


class Package:
    """An OOXML zip held in memory: parts are parsed on first use and written back only if
    marked changed."""

    def __init__(self, path):
        self.path = Path(path)
        with zipfile.ZipFile(self.path) as archive:
            self.infos = archive.infolist()
            self.data = {info.filename: archive.read(info) for info in self.infos}
        self.trees = {}
        self.changed = set()

    def names(self, pattern=None):
        found = [name for name in self.data if pattern is None or pattern.match(name)]
        return sorted(found, key=number_key)

    def xml(self, name):
        if name not in self.data:
            return None
        if name not in self.trees:
            self.trees[name] = etree.fromstring(self.data[name], PARSER)
        return self.trees[name]

    def put_xml(self, name, root):
        self.trees[name] = root
        self.changed.add(name)

    def put_bytes(self, name, data):
        self.data[name] = data
        self.trees.pop(name, None)
        self.changed.add(name)

    def mark(self, name):
        self.changed.add(name)

    def save(self, out):
        for name in self.changed:
            if name in self.trees:
                self.data[name] = etree.tostring(self.trees[name], xml_declaration=True, encoding='UTF-8', standalone=True)
        known = {info.filename for info in self.infos}
        with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as archive:
            for info in self.infos:
                archive.writestr(info, self.data[info.filename], compress_type=info.compress_type)
            for name in self.data:
                if name not in known:
                    archive.writestr(name, self.data[name])


class Run:
    """One run of text: where it is, what it says, and the font and sizes it resolves to."""

    def __init__(self, part, element, text, font, source, size=None, size_cs=None, bijoy_font=None):
        self.part = part
        self.element = element
        self.text = text
        self.font = font
        self.source = source
        """Where `font` came from: 'run', 'style', 'defaults', 'theme', or None when unset."""
        self.size = size
        self.size_cs = size_cs
        self.size_cs_behind = False
        """DOCX: the Latin size is set closer to the run than the complex-script one, so the
        Bangla is drawn at an inherited size nobody chose for it."""
        self.bijoy_font = bijoy_font if bijoy_font and not has_bangla(text) and text.strip() else None
        """The Bijoy font the run's Latin letters are drawn in: its text is Bijoy-encoded Bangla."""

    @property
    def bangla(self):
        return has_bangla(self.text)


def readable(pieces):
    """The text a reader sees in (text, bijoy font or None) pieces: Bijoy stretches converted to
    Unicode as a whole, since their reordering crosses run boundaries."""
    out, group = [], []
    for text, bijoy_font in pieces:
        if bijoy_font:
            group.append(text)
            continue
        if group:
            out.append(bijoy.to_unicode(''.join(group)))
            group = []
        out.append(text)
    if group:
        out.append(bijoy.to_unicode(''.join(group)))
    return ''.join(out)


def theme_fonts(root):
    """{'minor': font, 'major': font} a theme gives Bengali text: its Beng script font, else its
    complex-script font."""
    fonts = {}
    if root is None:
        return fonts
    for kind in ('minor', 'major'):
        scheme = root.find(f'.//a:fontScheme/a:{kind}Font', NS)
        if scheme is None:
            continue
        script = scheme.find("a:font[@script='Beng']", NS)
        cs = scheme.find('a:cs', NS)
        name = (script.get('typeface') if script is not None else '') or (cs.get('typeface') if cs is not None else '')
        if name:
            fonts[kind] = name
    return fonts


class Docx:
    """Resolves Word's formatting chain for the runs of a document."""

    def __init__(self, package):
        self.package = package
        self.styles = {}
        self.default_paragraph_style = None
        self.defaults = None
        styles = package.xml('word/styles.xml')
        if styles is not None:
            for style in styles.findall('w:style', NS):
                self.styles[style.get(w('styleId'))] = style
                if style.get(w('type')) == 'paragraph' and style.get(w('default')) in ('1', 'true', 'on'):
                    self.default_paragraph_style = style.get(w('styleId'))
            self.defaults = styles.find('w:docDefaults/w:rPrDefault/w:rPr', NS)
        theme = next((package.xml(name) for name in package.names(re.compile(r'^word/theme/theme\d*\.xml$'))), None)
        self.theme = theme_fonts(theme)

    def _chain(self, style_id):
        seen = []
        while style_id and style_id in self.styles and style_id not in seen and len(seen) < MAX_STYLE_DEPTH:
            seen.append(style_id)
            based = self.styles[style_id].find('w:basedOn', NS)
            style_id = based.get(w('val')) if based is not None else None
        return [self.styles[style_id] for style_id in seen]

    def _sources(self, run):
        """(label, rPr) in the order Word looks for a run property."""
        yield 'run', run.find('w:rPr', NS)
        rstyle = run.find('w:rPr/w:rStyle', NS)
        for style in self._chain(rstyle.get(w('val')) if rstyle is not None else None):
            yield 'style', style.find('w:rPr', NS)
        paragraph = next(run.iterancestors(w('p')), None)
        pstyle = paragraph.find('w:pPr/w:pStyle', NS) if paragraph is not None else None
        for style in self._chain(pstyle.get(w('val')) if pstyle is not None else self.default_paragraph_style):
            yield 'style', style.find('w:rPr', NS)
        yield 'defaults', self.defaults

    def cs_font(self, run):
        """(font, source) Word uses for Bangla in `run`."""
        for label, rpr in self._sources(run):
            fonts = rpr.find('w:rFonts', NS) if rpr is not None else None
            if fonts is None:
                continue
            theme = fonts.get(w('cstheme'))
            if theme:
                return self.theme.get('major' if theme.startswith('major') else 'minor'), 'theme'
            if fonts.get(w('cs')):
                return fonts.get(w('cs')), label
        return self.theme.get('minor'), 'theme' if self.theme.get('minor') else None

    def latin_font(self, run, slot):
        """The font Word names in `slot` ('ascii' or 'hAnsi') for `run`; None for a theme font."""
        for _, rpr in self._sources(run):
            fonts = rpr.find('w:rFonts', NS) if rpr is not None else None
            if fonts is None:
                continue
            if fonts.get(w(f'{slot}Theme')):
                return None
            if fonts.get(w(slot)):
                return fonts.get(w(slot))
        return None

    def bijoy_font(self, run):
        """The Bijoy font Word draws `run`'s Latin letters with, or None.

        Word draws ASCII with the w:ascii font and the Windows-1252 codes Bijoy also uses (†, ¨, ©)
        with w:hAnsi, so either slot naming SutonnyMJ makes the run Bijoy."""
        for slot in ('ascii', 'hAnsi'):
            font = self.latin_font(run, slot)
            if bijoy.is_font(font):
                return font
        return None

    def half_points(self, run, name):
        """(size in half points, how far up the chain it was found), or (None, None)."""
        for depth, (_, rpr) in enumerate(self._sources(run)):
            size = rpr.find(f'w:{name}', NS) if rpr is not None else None
            if size is not None and size.get(w('val')):
                return int(size.get(w('val'))), depth
        return None, None

    def runs(self):
        for part in self.package.names(DOCX_TEXT_PARTS):
            root = self.package.xml(part)
            for run in root.iter(w('r')):
                text = run_text(run)
                if not text:
                    continue
                font, source = self.cs_font(run)
                size, size_depth = self.half_points(run, 'sz')
                size_cs, size_cs_depth = self.half_points(run, 'szCs')
                points = (size or DEFAULT_HALF_POINTS) / 2
                points_cs = (size_cs or DEFAULT_HALF_POINTS) / 2
                found = Run(part, run, text, font, source, points, points_cs, self.bijoy_font(run))
                found.size_cs_behind = size_depth is not None and (size_cs_depth is None or size_cs_depth > size_depth)
                yield found

    def paragraphs(self):
        """(part, text) of every paragraph, text boxes counted as their own paragraphs, with Bijoy
        text converted to Unicode."""
        for part in self.package.names(DOCX_TEXT_PARTS):
            for paragraph in self.package.xml(part).iter(w('p')):
                pieces = [
                    (run_text(run), self.bijoy_font(run))
                    for run in paragraph.iter(w('r'))
                    if next(run.iterancestors(w('p')), None) is paragraph
                ]
                text = readable(pieces)
                if text.strip():
                    yield part, text

    def text_boxes(self):
        return sum(len(self.package.xml(part).findall('.//w:txbxContent', NS)) for part in self.package.names(DOCX_TEXT_PARTS))

    def pictures(self):
        return sum(len(self.package.xml(part).findall('.//a:blip', NS)) for part in self.package.names(DOCX_TEXT_PARTS))

    def embedded_fonts(self):
        table = self.package.xml('word/fontTable.xml')
        if table is None:
            return set()
        return {
            font.get(w('name'))
            for font in table.findall('w:font', NS)
            if font.find('w:embedRegular', NS) is not None
        }


def run_text(run):
    parts = []
    for child in run:
        if child.tag == w('t'):
            parts.append(child.text or '')
        elif child.tag == w('tab'):
            parts.append('\t')
    return ''.join(parts)


class Pptx:
    """Resolves PowerPoint's font chain for the runs of a deck, as far as the slide and theme."""

    def __init__(self, package):
        self.package = package
        theme = next((package.xml(name) for name in package.names(re.compile(r'^ppt/theme/theme\d*\.xml$'))), None)
        self.theme = theme_fonts(theme)

    def _theme_font(self, typeface):
        if typeface.startswith('+mj'):
            return self.theme.get('major'), 'theme'
        if typeface.startswith('+mn'):
            return self.theme.get('minor'), 'theme'
        return typeface, None

    def _typeface(self, run, tag):
        """(typeface, source) of the `a:<tag>` font set on the run, its paragraph or its text box's
        list style, or (None, None)."""
        paragraph = next(run.iterancestors(a('p')), None)
        level = 1
        ppr = paragraph.find('a:pPr', NS) if paragraph is not None else None
        if ppr is not None and ppr.get('lvl'):
            level = int(ppr.get('lvl')) + 1
        body = next(run.iterancestors(f'{{{P}}}txBody'), None)
        if body is None:
            body = next(run.iterancestors(a('txBody')), None)
        candidates = [
            ('run', run.find(f'a:rPr/a:{tag}', NS)),
            ('style', paragraph.find(f'a:pPr/a:defRPr/a:{tag}', NS) if paragraph is not None else None),
            ('style', body.find(f'a:lstStyle/a:lvl{level}pPr/a:defRPr/a:{tag}', NS) if body is not None else None),
        ]
        for label, font in candidates:
            if font is not None and font.get('typeface'):
                name, theme = self._theme_font(font.get('typeface'))
                return name, theme or label
        return None, None

    def cs_font(self, run):
        font, source = self._typeface(run, 'cs')
        if source:
            return font, source
        shape = next(run.iterancestors(f'{{{P}}}sp'), None)
        title = shape is not None and shape.find(".//p:nvPr/p:ph[@type='title']", NS) is not None
        title = title or (shape is not None and shape.find(".//p:nvPr/p:ph[@type='ctrTitle']", NS) is not None)
        kind = 'major' if title else 'minor'
        return self.theme.get(kind), 'theme' if self.theme.get(kind) else None

    def runs(self):
        for part in self.package.names(PPTX_TEXT_PARTS):
            for run in self.package.xml(part).iter(a('r')):
                t = run.find('a:t', NS)
                text = t.text if t is not None else ''
                if not text:
                    continue
                font, source = self.cs_font(run)
                rpr = run.find('a:rPr', NS)
                size = int(rpr.get('sz')) / 100 if rpr is not None and rpr.get('sz') else None
                yield Run(part, run, text, font, source, size, size, self.bijoy_font(run))

    def bijoy_font(self, run):
        """The Bijoy font PowerPoint draws `run`'s Latin letters with (`a:latin`), or None."""
        font, source = self._typeface(run, 'latin')
        return font if source != 'theme' and bijoy.is_font(font) else None

    def paragraphs(self):
        for part in self.package.names(PPTX_TEXT_PARTS):
            for paragraph in self.package.xml(part).iter(a('p')):
                pieces = []
                for run in paragraph.iter(a('r')):
                    t = run.find('a:t', NS)
                    pieces.append((t.text or '' if t is not None else '', self.bijoy_font(run)))
                text = readable(pieces)
                if text.strip():
                    yield part, text

    def text_boxes(self):
        return sum(
            1
            for part in self.package.names(PPTX_TEXT_PARTS)
            for body in self.package.xml(part).iter(f'{{{P}}}txBody')
            if ''.join(t.text or '' for t in body.iter(a('t'))).strip()
        )

    def pictures(self):
        return sum(len(self.package.xml(part).findall('.//a:blip', NS)) for part in self.package.names(PPTX_TEXT_PARTS))

    def embedded_fonts(self):
        presentation = self.package.xml('ppt/presentation.xml')
        if presentation is None:
            return set()
        return {font.get('typeface') for font in presentation.findall('p:embeddedFontLst/p:embeddedFont/p:font', NS)}


class Xlsx:
    """The text cells of a workbook: one run per cell, or per rich-text run, with its font."""

    def __init__(self, path):
        from openpyxl import load_workbook

        self.workbook = load_workbook(path, rich_text=True)

    def _cells(self):
        """(sheet title, cell) for every cell holding text rather than a number or a formula."""
        from openpyxl.cell.rich_text import CellRichText

        for sheet in self.workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    text = cell.value
                    if isinstance(text, CellRichText) or (isinstance(text, str) and not text.startswith('=')):
                        yield sheet.title, cell

    def _cell_runs(self, part, cell):
        """Runs of one cell: rich-text runs name their own font, or inherit the cell's."""
        pieces = [cell.value] if isinstance(cell.value, str) else list(cell.value)
        for piece in pieces:
            text = piece if isinstance(piece, str) else piece.text
            inline = getattr(getattr(piece, 'font', None), 'rFont', None)
            name = inline or cell.font.name
            size = cell.font.sz
            source = 'run' if name else None
            yield Run(part, cell, text or '', name, source, size, size, name if bijoy.is_font(name) else None)

    def runs(self):
        for part, cell in self._cells():
            yield from self._cell_runs(part, cell)

    def paragraphs(self):
        for part, cell in self._cells():
            pieces = [(run.text, run.bijoy_font) for run in self._cell_runs(part, cell)]
            yield f'{part}!{cell.coordinate}', readable(pieces)

    def text_boxes(self):
        return 0

    def pictures(self):
        return 0

    def embedded_fonts(self):
        return set()


def open_document(path):
    """(kind, reader) for a .docx, .pptx or .xlsx file."""
    suffix = Path(path).suffix.lower()
    if suffix in ('.docx', '.dotx', '.docm'):
        return 'docx', Docx(Package(path))
    if suffix in ('.pptx', '.potx', '.pptm'):
        return 'pptx', Pptx(Package(path))
    if suffix in ('.xlsx', '.xltx', '.xlsm'):
        return 'xlsx', Xlsx(path)
    raise ValueError(f'not a DOCX, PPTX or XLSX file: {path}')
