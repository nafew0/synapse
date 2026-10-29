"""Write values, formulas and styles straight into a workbook's XML, leaving everything else as it was.

    python3 cells.py <in.xlsx> <out.xlsx> --set "Budget!D11==C11/B11-1" --set "Budget!A12=Total"
                     [--style "Budget!D11:D12=Budget!D10"] [--width "Budget!E=14"]

--set REF=VALUE   a number, text, or a formula when VALUE starts with "=" ("=" alone clears)
--style RANGE=REF give every cell in RANGE the formatting of REF (font, fill, border, number format)
--width SHEET!COL=N set a column width

Use this instead of openpyxl for a workbook that preflight.py says openpyxl would damage
(sparklines, text boxes, rich text, Excel-only conditional formats). Cells and rows are
created in order when missing. Formulas get no cached value: run recalc.py afterwards.
"""

import argparse
import re
import shutil
import sys
import tempfile
import zipfile

from lxml import etree
from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries

from parts import MAIN_NS, sheet_parts

M = f"{{{MAIN_NS}}}"
REF = re.compile(r"^(?:'((?:[^']|'')+)'|([^!]+))!([A-Z]{1,3}\d+(?::[A-Z]{1,3}\d+)?)$")
CELL = re.compile(r"^([A-Z]{1,3})(\d+)$")
NUMBER = re.compile(r"^-?\d+(\.\d+)?([eE][-+]?\d+)?$")


class Usage(Exception):
    pass


def split_ref(text):
    match = REF.match(text.strip())
    if not match:
        raise Usage(f"expected Sheet!A1, got {text!r}")
    sheet = (match.group(1) or "").replace("''", "'") or match.group(2)
    return sheet, match.group(3)


def cell_order(ref):
    letters, row = CELL.match(ref).groups()
    return int(row), column_index_from_string(letters)


class Sheet:
    def __init__(self, xml):
        self.root = etree.fromstring(xml)
        self.data = self.root.find(f"{M}sheetData")

    def row(self, number):
        rows = self.data.findall(f"{M}row")
        for row in rows:
            current = int(row.get("r"))
            if current == number:
                return row
            if current > number:
                new = etree.Element(f"{M}row", r=str(number))
                row.addprevious(new)
                return new
        return etree.SubElement(self.data, f"{M}row", r=str(number))

    def cell(self, ref):
        number, column = cell_order(ref)
        row = self.row(number)
        for cell in row.findall(f"{M}c"):
            current = cell_order(cell.get("r"))[1]
            if current == column:
                return cell
            if current > column:
                new = etree.Element(f"{M}c", r=ref)
                cell.addprevious(new)
                return new
        return etree.SubElement(row, f"{M}c", r=ref)

    def find(self, ref):
        for cell in self.data.iter(f"{M}c"):
            if cell.get("r") == ref:
                return cell
        return None

    def widen_dimension(self, ref):
        dimension = self.root.find(f"{M}dimension")
        if dimension is None:
            return
        current = dimension.get("ref", "A1")
        first, _, last = current.partition(":")
        last = last or first
        rows = [cell_order(first)[0], cell_order(last)[0], cell_order(ref)[0]]
        cols = [cell_order(first)[1], cell_order(last)[1], cell_order(ref)[1]]
        dimension.set("ref", f"{get_column_letter(min(cols))}{min(rows)}:{get_column_letter(max(cols))}{max(rows)}")


def set_value(cell, value):
    for child in list(cell):
        cell.remove(child)
    cell.attrib.pop("t", None)
    if value == "=":
        return
    if value.startswith("="):
        etree.SubElement(cell, f"{M}f").text = value[1:]
    elif NUMBER.match(value):
        etree.SubElement(cell, f"{M}v").text = value
    elif value.upper() in ("TRUE", "FALSE"):
        cell.set("t", "b")
        etree.SubElement(cell, f"{M}v").text = "1" if value.upper() == "TRUE" else "0"
    else:
        cell.set("t", "inlineStr")
        text = etree.SubElement(etree.SubElement(cell, f"{M}is"), f"{M}t")
        text.text = value
        if value != value.strip():
            text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def set_width(sheet, column, width):
    index = column_index_from_string(column)
    cols = sheet.root.find(f"{M}cols")
    if cols is None:
        cols = etree.Element(f"{M}cols")
        sheet.root.find(f"{M}sheetData").addprevious(cols)
    for col in cols.findall(f"{M}col"):
        low, high = int(col.get("min")), int(col.get("max"))
        if low <= index <= high:
            if low != high:
                raise Usage(f"column {column} shares a width setting with its neighbours; set the whole range")
            col.set("width", str(width))
            col.set("customWidth", "1")
            return
    new = etree.Element(f"{M}col", min=str(index), max=str(index), width=str(width), customWidth="1")
    later = [c for c in cols.findall(f"{M}col") if int(c.get("min")) > index]
    (later[0].addprevious(new) if later else cols.append(new))


def main(argv):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("source")
    parser.add_argument("output")
    parser.add_argument("--set", action="append", default=[], dest="values")
    parser.add_argument("--style", action="append", default=[])
    parser.add_argument("--width", action="append", default=[])
    args = parser.parse_args(argv)

    with zipfile.ZipFile(args.source) as archive:
        parts = sheet_parts(archive)
        sheets = {}

        def load(name):
            if name not in parts:
                raise Usage(f"no sheet named {name!r}; sheets are {list(parts)}")
            if name not in sheets:
                sheets[name] = Sheet(archive.read(parts[name]))
            return sheets[name]

        changed = 0
        for spec in args.values:
            target, _, value = spec.partition("=")
            name, ref = split_ref(target)
            sheet = load(name)
            set_value(sheet.cell(ref), value)
            sheet.widen_dimension(ref)
            changed += 1
        for spec in args.style:
            target, _, source = spec.partition("=")
            name, cells = split_ref(target)
            source_name, source_ref = split_ref(source)
            source_cell = load(source_name).find(source_ref)
            style = source_cell.get("s") if source_cell is not None else None
            min_col, min_row, max_col, max_row = range_boundaries(cells)
            sheet = load(name)
            for row in range(min_row, max_row + 1):
                for col in range(min_col, max_col + 1):
                    cell = sheet.cell(f"{get_column_letter(col)}{row}")
                    if style is None:
                        cell.attrib.pop("s", None)
                    else:
                        cell.set("s", style)
                    sheet.widen_dimension(cell.get("r"))
                    changed += 1
        for spec in args.width:
            target, _, width = spec.partition("=")
            name, _, column = target.rpartition("!")
            set_width(load(name.strip("'")), column, float(width))

        replaced = {parts[name]: etree.tostring(sheet.root, xml_declaration=True, encoding="UTF-8", standalone=True) for name, sheet in sheets.items()}
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as handle:
            temporary = handle.name
        with zipfile.ZipFile(temporary, "w") as out:
            for info in archive.infolist():
                data = replaced.get(info.filename)
                out.writestr(info, data if data is not None else archive.read(info.filename))
    shutil.move(temporary, args.output)
    print(f"wrote {changed} cell change(s) to {args.output}; every other part is unchanged. Run recalc.py next if you set formulas.")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main(sys.argv[1:]))
    except Usage as error:
        print(f"error: {error}")
        sys.exit(1)
