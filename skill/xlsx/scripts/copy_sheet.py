#!/usr/bin/env python3
"""Add a tab that is a faithful copy of an existing sheet, ready to fill in.

openpyxl's copy_worksheet copies values, styles, merges, column widths, row heights, margins and
page setup, but silently leaves out the page header and footer, print titles, print area, page
breaks, freeze panes, zoom, gridlines, protection, data validation and conditional formatting.
This copies those too and puts the new tab right after its template.

Usage:
    python copy_sheet.py ORIGINAL OUTPUT --template 'Sheet1' --name 'Investigation Committee'

Prints JSON and exits 0 on success, 1 on bad input. Images and charts cannot be copied by
openpyxl; they are listed under `not_copied` so you can tell the user.
"""
import argparse
import json
import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

from layout import header_footer, print_area

INVALID_SHEET_CHARACTERS = set(':\\/?*[]')
MAX_SHEET_NAME = 31


class InputError(Exception):
    pass


def check_name(name, existing):
    if not name or len(name) > MAX_SHEET_NAME or INVALID_SHEET_CHARACTERS & set(name):
        raise InputError(f'Invalid sheet name {name!r}: 1-31 characters, none of : \\ / ? * [ ]')
    if name in existing:
        raise InputError(f'A sheet named {name!r} already exists')


def copy_print_layout(template, ws):
    """Header and footer, print titles and area, page breaks."""
    copied = []
    ws.HeaderFooter = copy(template.HeaderFooter)
    if header_footer(template):
        copied.append('header/footer')
    if template.print_title_rows:
        ws.print_title_rows = template.print_title_rows.replace('$', '')
        copied.append('print title rows')
    if template.print_title_cols:
        ws.print_title_cols = template.print_title_cols.replace('$', '')
        copied.append('print title columns')
    area = print_area(template)
    if area:
        ws.print_area = area.split(',')
        copied.append('print area')
    for kind in ('row_breaks', 'col_breaks'):
        for brk in getattr(template, kind).brk:
            getattr(ws, kind).append(copy(brk))
    if template.row_breaks.brk or template.col_breaks.brk:
        copied.append('page breaks')
    return copied


def copy_view_and_rules(template, ws):
    """Freeze panes, zoom, gridlines, protection, data validation, conditional formatting."""
    copied = []
    ws.freeze_panes = template.freeze_panes
    view, source = ws.sheet_view, template.sheet_view
    view.zoomScale, view.showGridLines, view.view = source.zoomScale, source.showGridLines, source.view
    ws.protection = copy(template.protection)
    for validation in template.data_validations.dataValidation:
        ws.add_data_validation(copy(validation))
    if template.data_validations.dataValidation:
        copied.append('data validation')
    for formatting in template.conditional_formatting:
        for rule in formatting.rules:
            ws.conditional_formatting.add(str(formatting.sqref), copy(rule))
    if len(template.conditional_formatting):
        copied.append('conditional formatting')
    return copied


def copy_sheet(original, output, template_name, name):
    wb = load_workbook(original)
    if template_name not in wb.sheetnames:
        raise InputError(f'No sheet named {template_name!r}; sheets are {wb.sheetnames}')
    check_name(name, wb.sheetnames)
    template = wb[template_name]
    ws = wb.copy_worksheet(template)
    ws.title = name
    copied = copy_print_layout(template, ws) + copy_view_and_rules(template, ws)
    wb.move_sheet(ws, offset=wb.sheetnames.index(template_name) + 1 - wb.sheetnames.index(name))
    wb.save(output)
    not_copied = [
        f'{count} {kind}' for kind, count in (('image(s)', len(template._images)), ('chart(s)', len(template._charts))) if count
    ]
    return {
        'status': 'success',
        'output': str(output),
        'sheet': name,
        'template': template_name,
        'sheets': wb.sheetnames,
        'also_copied': copied,
        'not_copied': not_copied,
        'verify_with': f"--new-sheet '{name}:{template_name}'",
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('original', type=Path, help="the user's workbook, left untouched")
    parser.add_argument('output', type=Path, help='where to save the workbook with the new tab')
    parser.add_argument('--template', required=True, help='the sheet to copy')
    parser.add_argument('--name', required=True, help='the new tab name')
    args = parser.parse_args()

    try:
        if not args.original.is_file():
            raise InputError(f'File not found: {args.original}')
        if args.output.exists() and args.original.resolve() == args.output.resolve():
            raise InputError('ORIGINAL and OUTPUT are the same file; save the copy under a new name')
        result = copy_sheet(args.original, args.output, args.template, args.name)
    except InputError as error:
        print(json.dumps({'error': str(error)}, indent=2))
        return 1

    print(json.dumps(result, indent=2))
    return 0


if __name__ == '__main__':
    sys.exit(main())
