#!/usr/bin/env python3
"""Build a Bangladesh government document (DOCX or PDF) from a JSON spec, then check it.

The layout follows the ক্রোড়পত্র formats of সচিবালয় নির্দেশমালা ২০২৪: centred letterhead, number
on the left and the date pair (বঙ্গাব্দ above খ্রিষ্টাব্দ) on the right, numbered paragraphs,
signature block on the right, recipient and copy list where the type puts them. See formats.md.

Usage:
    python3 build.py SPEC.json OUTPUT.docx|OUTPUT.pdf [--no-verify]

After writing the DOCX it runs, in order, and stops at the first failure:
  1. the docx skill's fix_bangla.py --new (Nikosh in every slot, embedded)
  2. check.py (format, dates, spelling, digits)
  3. the docx skill's verify_bangla.py (LibreOffice render + OCR), unless --no-verify
For a .pdf output the checked DOCX is converted with LibreOffice; the DOCX stays in /mnt/data/qa.

Needs the docx skill in the sandbox: invoke it before running this script.
Prints JSON and exits 0 when everything passed, 2 when a check found defects, 1 on bad input.
"""
import argparse
import json
import re
import shutil
import subprocess
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))

import bn  # noqa: E402
import check  # noqa: E402

SKILLS = Path(__file__).resolve().parents[2]
OFFICE = SKILLS / 'docx' / 'scripts' / 'office'
QA = Path('/mnt/data/qa')

FONT = 'Nikosh'
LATIN_FONT = 'Times New Roman'
BODY_PT = 13
LETTERHEAD_PT = 15
FOOTER_PT = 11
TEXT_WIDTH_CM = 16.0
INDENT = Cm(1.27)
ITEM_INDENT = Cm(2.54)
HANG = Cm(-1.27)
LIST_HANG = Cm(-0.9)
SIGNATURE_INDENT = Cm(8.5)
RPR_ORDER = ('w:rStyle', 'w:rFonts', 'w:b', 'w:bCs', 'w:i', 'w:iCs', 'w:caps', 'w:smallCaps', 'w:strike',
             'w:dstrike', 'w:outline', 'w:shadow', 'w:emboss', 'w:imprint', 'w:noProof', 'w:snapToGrid',
             'w:vanish', 'w:webHidden', 'w:color', 'w:spacing', 'w:w', 'w:kern', 'w:position', 'w:sz',
             'w:szCs', 'w:highlight', 'w:u', 'w:effect', 'w:bdr', 'w:shd', 'w:fitText', 'w:vertAlign',
             'w:rtl', 'w:cs', 'w:em', 'w:lang', 'w:eastAsianLayout', 'w:specVanish', 'w:oMath')
ITEM_LETTERS = 'কখগঘঙচছজঝঞটঠডঢ'
LATIN_TOKEN = re.compile(r'[A-Za-z@/\\]')
DIGIT_TOKEN = re.compile(r'\S*[0-9]\S*')

COPIES = 'অবগতি ও প্রয়োজনীয় ব্যবস্থা গ্রহণের জন্য অনুলিপি প্রেরণ করা হলো:'

TYPES = {
    'letter': dict(label='নম্বর:', to='top', subject='required', salutation='মহোদয়,', closing='আপনার বিশ্বস্ত,'),
    'memo': dict(label='স্মারক নম্বর:', to='bottom', subject='required'),
    'order': dict(title='অফিস আদেশ', label='নম্বর:', subject='none'),
    'circular': dict(title='পরিপত্র', label='নম্বর:', subject='required', distribution='বিতরণ:'),
    'notification': dict(title='প্রজ্ঞাপন', label='নম্বর:', authority='রাষ্ট্রপতির আদেশক্রমে,', subject='none',
                         copy_label='স্মারক নম্বর:'),
    'notice': dict(title='বিজ্ঞপ্তি', label='নম্বর:', subject='required',
                   distribution='বিতরণ (জ্ঞাতার্থে ও কার্যার্থে):'),
    'meeting': dict(title='সভার নোটিশ', label='নম্বর:', subject='required', distribution='বিতরণ:'),
    'do': dict(label='আধা-সরকারি পত্র নম্বর:', to='bottom', salutation='প্রিয় মহোদয়,',
               closing='আন্তরিকভাবে আপনার,', regards='শুভেচ্ছান্তে,', letterhead=False, sender=True, subject='none'),
    'noc': dict(title='অনাপত্তি সনদ', label='নম্বর:', numbered=False),
    'certificate': dict(title='প্রত্যয়নপত্র', label='নম্বর:', numbered=False),
    'minutes': dict(title='সভার কার্যবিবরণী', label='নম্বর:', numbered=False),
}


class SpecError(ValueError):
    pass


def bangla_digits(text):
    """Latin digits as Bangla, except inside e-mail addresses, URLs and Latin words."""
    return DIGIT_TOKEN.sub(lambda m: m.group(0) if LATIN_TOKEN.search(m.group(0)) else bn.digits(m.group(0)), text)


def child(rpr, tag):
    """The `tag` child of a w:rPr, created in schema order when missing."""
    element = rpr.find(qn(tag))
    if element is None:
        element = OxmlElement(tag)
        rpr.insert_element_before(element, *RPR_ORDER[RPR_ORDER.index(tag) + 1:])
    return element


def set_fonts(rpr, size_pt=None, bold=None):
    fonts = child(rpr, 'w:rFonts')
    for slot in ('w:ascii', 'w:hAnsi'):
        fonts.set(qn(slot), LATIN_FONT)
    for slot in ('w:eastAsia', 'w:cs'):
        fonts.set(qn(slot), FONT)
    for theme in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        fonts.attrib.pop(qn(theme), None)
    if bold is not None:
        for tag in ('w:b', 'w:bCs'):
            if bold:
                child(rpr, tag)
            elif rpr.find(qn(tag)) is not None:
                rpr.remove(rpr.find(qn(tag)))
    if size_pt is not None:
        for tag in ('w:sz', 'w:szCs'):
            child(rpr, tag).set(qn('w:val'), str(round(size_pt * 2)))
    language = child(rpr, 'w:lang')
    language.set(qn('w:val'), 'bn-BD')
    language.set(qn('w:bidi'), 'bn-BD')


def setup(document):
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.75)
    section.right_margin = Cm(2.25)
    defaults = document.styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
    set_fonts(defaults, BODY_PT)
    normal = document.styles['Normal']
    set_fonts(normal.element.get_or_add_rPr(), BODY_PT)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.15
    footer = document.styles['Footer']
    set_fonts(footer.element.get_or_add_rPr(), FOOTER_PT)
    properties = document.core_properties
    properties.author = properties.last_modified_by = 'Synapse'


def paragraph(container, text='', align=None, bold=False, size=None, indent=None, first=None,
              before=0, after=0, underline=False, keep=False):
    p = container.add_paragraph()
    if align is not None:
        p.alignment = align
    fmt = p.paragraph_format
    fmt.space_before, fmt.space_after = Pt(before), Pt(after)
    if indent is not None:
        fmt.left_indent = indent
    if first is not None:
        fmt.first_line_indent = first
    if keep:
        fmt.keep_with_next = True
    if text:
        run = p.add_run(bangla_digits(text))
        run.bold = bold or None
        run.underline = underline or None
        if bold or size:
            set_fonts(run._element.get_or_add_rPr(), size, bold)
    return p


def bottom_border(p):
    borders = OxmlElement('w:pBdr')
    line = OxmlElement('w:bottom')
    for key, value in (('w:val', 'single'), ('w:sz', '4'), ('w:space', '1'), ('w:color', 'auto')):
        line.set(qn(key), value)
    borders.append(line)
    p._p.get_or_add_pPr().append(borders)


def grid(document, widths_cm, rows=1, style=None):
    table = document.add_table(rows=rows, cols=len(widths_cm))
    if style:
        table.style = document.styles[style]
    table.autofit = False
    for column, width in zip(table.columns, widths_cm):
        column.width = Cm(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)
    return table


def cell_text(cell, lines, align=None, bold=False):
    first = cell.paragraphs[0]
    for i, line in enumerate(lines):
        p = first if i == 0 else cell.add_paragraph()
        if align is not None:
            p.alignment = align
        if line:
            run = p.add_run(bangla_digits(line))
            if bold:
                run.bold = True
                set_fonts(run._element.get_or_add_rPr(), None, True)
    return cell


def dates(spec):
    value = spec.get('date', 'today')
    if not value:
        return bn.PLACEHOLDER[:12], bn.PLACEHOLDER[:12]
    day = date.today() if value == 'today' else bn.parse_date(value)
    suffix = bool(spec.get('date_suffix'))
    return bn.bangla_date(day, suffix), bn.english_date(day, suffix)


def number_line(document, label, number, spec):
    """নম্বর on the left; তারিখ with বঙ্গাব্দ above খ্রিষ্টাব্দ on the right (ক্রোড়পত্র-১৭)."""
    bangla, english = dates(spec)
    table = grid(document, (10.5, 1.7, 3.8))
    left, middle, right = table.rows[0].cells
    cell_text(left, [label.replace(' ', '\u00a0') + '\u00a0' + number])
    cell_text(middle, ['তারিখ:'], WD_ALIGN_PARAGRAPH.RIGHT)
    cell_text(right, [bangla, english], WD_ALIGN_PARAGRAPH.CENTER)
    bottom_border(right.paragraphs[0])
    for cell in (left, middle):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    paragraph(document, after=6)


def letterhead(document, lines):
    for i, line in enumerate(lines):
        paragraph(document, line, WD_ALIGN_PARAGRAPH.CENTER, bold=i == 0, size=LETTERHEAD_PT if i == 0 else None)
    paragraph(document, after=4)


def sender_block(document, sender):
    lines = [sender.get('name'), sender.get('designation'), sender.get('office'),
             _labelled('ফোন', sender.get('phone')), _labelled('মোবাইল', sender.get('mobile')),
             _labelled('ই-মেইল', sender.get('email'))]
    for line in filter(None, lines):
        paragraph(document, line, indent=SIGNATURE_INDENT)
    paragraph(document, after=6)


def _labelled(label, value):
    return f'{label}: {value}' if value else None


def body(document, blocks, numbered):
    count = 0
    for block in blocks:
        if isinstance(block, str):
            count += 1
            prefix = f'{bn.digits(count)}। ' if numbered and count > 1 else ''
            paragraph(document, prefix + block, WD_ALIGN_PARAGRAPH.JUSTIFY, first=INDENT, after=6)
        elif 'items' in block:
            for i, item in enumerate(block['items']):
                text = item if item.startswith('(') else f'({ITEM_LETTERS[i % len(ITEM_LETTERS)]}) {item}'
                paragraph(document, text, WD_ALIGN_PARAGRAPH.JUSTIFY, indent=ITEM_INDENT, first=HANG, after=3)
            paragraph(document, after=3)
        elif 'table' in block:
            table_block(document, block['table'])
        else:
            raise SpecError(f'unknown body block: {block}')


def table_block(document, rows):
    if not rows:
        return
    columns = len(rows[0])
    table = grid(document, [TEXT_WIDTH_CM / columns] * columns, rows=len(rows), style='Table Grid')
    for r, row in enumerate(rows):
        for cell, value in zip(table.rows[r].cells, row):
            cell_text(cell, str(value).split('\n'), bold=r == 0)
    paragraph(document, after=6)


def signature(document, signer, authority=None, closing=None, regards=None):
    if regards:
        paragraph(document, regards, before=12)
    for line in filter(None, (closing, authority)):
        paragraph(document, line, WD_ALIGN_PARAGRAPH.CENTER, indent=SIGNATURE_INDENT, before=6, keep=True)
    paragraph(document, before=24, keep=True)
    name = signer.get('name')
    lines = [f'({name})' if name else None, _labelled('পরিচিতি নম্বর', signer.get('id')),
             signer.get('designation'), signer.get('office'), _labelled('ফোন', signer.get('phone')),
             _labelled('ই-মেইল', signer.get('email'))]
    for line in filter(None, lines):
        paragraph(document, line, WD_ALIGN_PARAGRAPH.CENTER, indent=SIGNATURE_INDENT, keep=True)
    paragraph(document, after=6)


def listing(document, heading, entries):
    if not entries:
        return
    paragraph(document, heading, underline=True, before=6, after=3, keep=True)
    for i, entry in enumerate(entries, 1):
        paragraph(document, f'{bn.digits(i)}। {entry}', indent=INDENT, first=LIST_HANG)


def build(spec, output):
    kind = spec.get('type')
    if kind not in TYPES:
        raise SpecError(f'type must be one of {", ".join(TYPES)}')
    rules = TYPES[kind]
    signer = spec.get('signer') or {}
    document = Document()
    setup(document)

    if spec.get('priority'):
        paragraph(document, spec['priority'], WD_ALIGN_PARAGRAPH.RIGHT, bold=True)
    if rules.get('sender'):
        sender_block(document, spec.get('sender') or signer)
    if rules.get('letterhead', True):
        lines = spec.get('letterhead') or []
        if not lines:
            raise SpecError('letterhead is required for this type')
        letterhead(document, lines)

    number = bn.file_number(spec.get('number'), spec.get('issue'))
    number_line(document, spec.get('number_label') or rules['label'], number, spec)

    if rules.get('to') == 'top' and spec.get('to'):
        lines = spec['to']
        paragraph(document, f'প্রাপক: {lines[0]}')
        for line in lines[1:]:
            paragraph(document, line, indent=Cm(1.45))
        paragraph(document, after=4)
    if rules.get('title'):
        paragraph(document, spec.get('title') or rules['title'], WD_ALIGN_PARAGRAPH.CENTER, bold=True,
                  underline=True, after=6)
    for label, value in spec.get('meta') or []:
        paragraph(document, f'{label}: {value}')
    subject_rule = rules.get('subject', 'optional')
    if spec.get('subject') and subject_rule == 'none':
        raise SpecError(f'{kind} has no বিষয় line: put the matter in the body')
    if subject_rule == 'required' and not spec.get('subject'):
        raise SpecError(f'subject (বিষয়) is required for {kind}')
    if spec.get('subject'):
        paragraph(document, f'বিষয়: {spec["subject"]}', WD_ALIGN_PARAGRAPH.JUSTIFY, bold=True,
                  indent=Cm(1.3), first=Cm(-1.3), after=3)
    references = spec.get('references') or []
    if len(references) == 1:
        paragraph(document, f'সূত্র: {references[0]}', WD_ALIGN_PARAGRAPH.JUSTIFY, indent=Cm(1.3), first=Cm(-1.3))
    elif references:
        paragraph(document, 'সূত্র:')
        for i, reference in enumerate(references, 1):
            paragraph(document, f'({bn.digits(i)}) {reference}', indent=Cm(1.3))
    paragraph(document, after=4)

    salutation = spec.get('salutation', rules.get('salutation'))
    if salutation:
        paragraph(document, salutation, after=4)
    body(document, spec.get('body') or [], spec.get('numbered', rules.get('numbered', True)))

    attachments = spec.get('attachments') or []
    if attachments:
        text = attachments[0] if len(attachments) == 1 else '; '.join(
            f'({bn.digits(i)}) {a}' for i, a in enumerate(attachments, 1))
        paragraph(document, f'সংযুক্তি: {text}', before=6)

    signature(document, signer, spec.get('authority', rules.get('authority')),
              spec.get('closing', rules.get('closing')), spec.get('regards', rules.get('regards')))

    if rules.get('to') == 'bottom' and spec.get('to'):
        for line in spec['to']:
            paragraph(document, line)
        paragraph(document, after=6)

    listing(document, spec.get('distribution_heading') or rules.get('distribution', 'বিতরণ:'),
            spec.get('distribution'))
    if spec.get('copies'):
        copy_number = spec.get('copy_number') or number
        paragraph(document, before=6)
        number_line(document, rules.get('copy_label', rules['label']), bn.digits(copy_number), spec)
        listing(document, spec.get('copies_heading') or COPIES, spec['copies'])
        if signer.get('name'):
            paragraph(document, before=24)
            for line in filter(None, (f'({signer["name"]})', signer.get('designation'))):
                paragraph(document, line, WD_ALIGN_PARAGRAPH.CENTER, indent=SIGNATURE_INDENT)

    if spec.get('footer'):
        footer = document.sections[0].footer.paragraphs[0]
        footer.text = spec['footer']
        footer.style = document.styles['Footer']
    document.save(output)


def run_json(command):
    completed = subprocess.run(command, capture_output=True, text=True)
    try:
        return completed.returncode, json.loads(completed.stdout)
    except json.JSONDecodeError:
        return completed.returncode, {'error': (completed.stdout + completed.stderr).strip()[-2000:]}


def to_pdf(docx_path, pdf_path):
    QA.mkdir(parents=True, exist_ok=True)
    command = [sys.executable, str(OFFICE / 'soffice.py'), '--headless', '--convert-to', 'pdf',
               '--outdir', str(QA), str(docx_path)]
    subprocess.run(command, capture_output=True, text=True, check=True)
    shutil.move(str(QA / (docx_path.stem + '.pdf')), pdf_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('spec', type=Path)
    parser.add_argument('output', type=Path)
    parser.add_argument('--no-verify', action='store_true', help='skip the LibreOffice + OCR render check')
    args = parser.parse_args()
    if not (OFFICE / 'fix_bangla.py').is_file():
        print(json.dumps({'error': 'the docx skill is not loaded: invoke the docx skill, then run this again'}))
        return 1
    try:
        spec = json.loads(args.spec.read_text(encoding='utf-8'))
        pdf = args.output.suffix.lower() == '.pdf'
        docx_path = QA / (args.output.stem + '.docx') if pdf else args.output
        docx_path.parent.mkdir(parents=True, exist_ok=True)
        build(spec, docx_path)
    except (OSError, ValueError, KeyError, TypeError) as error:
        print(json.dumps({'error': f'{type(error).__name__}: {error}'}, ensure_ascii=False))
        return 1

    result = {'output': str(args.output), 'type': spec['type']}
    code, result['fix'] = run_json([sys.executable, str(OFFICE / 'fix_bangla.py'), str(docx_path), '--new'])
    if code:
        result['status'] = 'failed'
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 1
    result['check'] = check.check(docx_path, spec['type'])
    if result['check']['status'] != 'clean':
        result['status'] = 'defects_found'
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 2
    if not args.no_verify:
        code, result['verify'] = run_json([sys.executable, str(OFFICE / 'verify_bangla.py'), str(docx_path)])
        if code:
            result['status'] = 'defects_found'
            print(json.dumps(result, ensure_ascii=False, indent=2))
            return 2
    if pdf:
        to_pdf(docx_path, args.output)
    result['status'] = 'clean'
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == '__main__':
    sys.exit(main())
