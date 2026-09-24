"""Tests for the Bangla width estimate in skill/xlsx/scripts/fit.py.

Run with openpyxl, HarfBuzz and the sandbox's Bangla fonts installed; the measured tests skip
without them:
    python -m pytest tests/skill/xlsx
"""
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

SCRIPTS = Path(__file__).resolve().parents[3] / 'skill' / 'xlsx' / 'scripts'
sys.path.insert(0, str(SCRIPTS))
import fit  # noqa: E402

BANGLA = 'তদন্ত কমিটির সম্মানী বিবরণী'
HAS_NIKOSH = fit.measure('ক', 'Nikosh', 11) is not None and fit.can_draw('Nikosh')
measured = pytest.mark.skipif(not HAS_NIKOSH, reason='needs HarfBuzz and the Nikosh font')


def problems(text, width, font='Nikosh', size=12):
    wb = Workbook()
    ws = wb.active
    ws['A1'] = text
    ws['A1'].font = Font(name=font, size=size)
    ws['B1'] = 'next cell is full, so nothing spills'
    ws.column_dimensions['A'].width = width
    return fit.fit_problems(ws, ws, fit.base_font_size(wb), fit.base_font_name(wb))


def test_latin_text_is_estimated_as_before():
    assert fit.text_width('Total', 11, False, 11) == fit.text_width('Total', 11, False, 11, 'Nikosh', 'Calibri')


def test_bangla_without_a_shaper_uses_the_wide_estimate(monkeypatch):
    monkeypatch.setattr(fit, 'measure', lambda *args, **kwargs: None)
    width = fit.text_width(BANGLA, 11, False, 11, 'Nikosh', 'Calibri')
    expected = len(BANGLA.replace(' ', '')) * 11 * fit.BANGLA_CHAR_EM / (11 * fit.DIGIT_EM) * fit.BANGLA_MARGIN
    latin_spaces = BANGLA.count(' ') * fit.char_width(' ')
    assert width == pytest.approx(expected + latin_spaces, rel=0.05)


@measured
def test_bangla_gets_thirty_percent_spare_room():
    shaped = fit.measure(BANGLA, 'Nikosh', 12) / fit.digit_points('Calibri', 11)
    tight = shaped * 1.05
    roomy = shaped * fit.BANGLA_MARGIN * 1.02
    assert problems(BANGLA, tight / fit.TEXT_TOLERANCE)
    assert not problems(BANGLA, roomy / fit.TEXT_TOLERANCE)


@measured
def test_a_latin_font_is_measured_as_the_windows_fallback():
    in_arial = fit.text_width(BANGLA, 12, False, 11, 'Arial', 'Calibri')
    in_nikosh = fit.text_width(BANGLA, 12, False, 11, 'Nikosh', 'Calibri')
    assert in_arial > in_nikosh * 1.2
