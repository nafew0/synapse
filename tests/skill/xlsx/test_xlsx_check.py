"""Tests for skill/xlsx/scripts/check.py. Needs LibreOffice: run in the sandbox image."""
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

SCRIPT = Path(__file__).resolve().parents[3] / 'skill' / 'xlsx' / 'scripts' / 'check.py'
pytestmark = pytest.mark.skipif(not shutil.which('soffice'), reason='needs LibreOffice')


def budget(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Budget'
    for row, (name, last, this) in enumerate([('ICT', 210, 260), ('Audit', 45, 50)], 2):
        for column, value in enumerate((name, last, this), 1):
            sheet.cell(row=row, column=column, value=value).font = Font(name='Arial', size=11)
    workbook.save(path)


def add_total(source, target, total='=SUM(B2:B3)'):
    workbook = load_workbook(source)
    workbook['Budget']['B4'] = total
    workbook.save(target)


def check(tmp_path, *flags):
    return subprocess.run(
        [sys.executable, SCRIPT, 'edited.xlsx', '--original', 'budget.xlsx', *flags],
        cwd=tmp_path, capture_output=True, text=True,
    )


def test_a_correct_edit_passes_every_gate(tmp_path):
    budget(tmp_path / 'budget.xlsx')
    add_total(tmp_path / 'budget.xlsx', tmp_path / 'edited.xlsx')

    result = check(tmp_path, '--allow', 'Budget!B4', '--formula', 'Budget!B4', '--expect', 'Budget!B4=255')

    assert result.returncode == 0, result.stdout + result.stderr
    assert 'structure: success' in result.stdout
    assert 'calculations: success' in result.stdout
    assert (tmp_path / 'qa' / 'edited-pages.jpg').exists()


def test_a_hardcoded_total_outside_the_allowed_range_fails(tmp_path):
    budget(tmp_path / 'budget.xlsx')
    add_total(tmp_path / 'budget.xlsx', tmp_path / 'edited.xlsx', total=255)

    result = check(tmp_path, '--allow', 'Budget!C4', '--formula', 'Budget!B4')

    assert result.returncode == 2
    assert 'calculations: FAILED' in result.stdout
    assert 'structure: FAILED' in result.stdout
