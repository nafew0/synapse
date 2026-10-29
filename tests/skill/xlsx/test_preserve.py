"""An edit to a user's workbook must keep the design openpyxl and LibreOffice drop on save.

Covers skill/xlsx/scripts/preflight.py, cells.py, cache.py (through recalc.py) and the design
counts in verify_structure.py. The recalc test needs LibreOffice: run in the sandbox image.
"""
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Font

SCRIPTS = Path(__file__).resolve().parents[3] / 'skill' / 'xlsx' / 'scripts'
sys.path.insert(0, str(SCRIPTS))
import verify_structure  # noqa: E402

X14 = 'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main"'
SPARKLINE = (
    '<extLst><ext uri="{05C60535-1F16-4fd2-B633-F4F36F0B64E0}" ' + X14 + '><x14:sparklineGroups>'
    '<x14:sparklineGroup><x14:colorSeries rgb="FF376092"/><x14:sparklines><x14:sparkline>'
    '<xm:f>Budget!B2:C2</xm:f><xm:sqref>E2</xm:sqref></x14:sparkline></x14:sparklines>'
    '</x14:sparklineGroup></x14:sparklineGroups></ext></extLst>'
)


def run(script, *args, cwd):
    return subprocess.run([sys.executable, SCRIPTS / script, *map(str, args)], cwd=cwd, capture_output=True, text=True)


def workbook(path, sparkline=True):
    """A small budget with a rich-text note and, by default, a sparkline openpyxl cannot keep."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Budget'
    for row, values in enumerate([('Dept', 2025, 2026), ('ICT', 210, 260), ('Audit', 45, 50)], 1):
        for col, value in enumerate(values, 1):
            ws.cell(row=row, column=col, value=value).font = Font(name='Arial', size=11)
    ws['B2'].number_format = '_(* #,##0_);_(* (#,##0);_(* "-"_);_(@_)'
    ws['A5'] = CellRichText([TextBlock(InlineFont(b=True, color='FFC00000'), 'Draft: '), 'figures in lakh'])
    wb.save(path)
    if not sparkline:
        return path
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    sheet = parts['xl/worksheets/sheet1.xml'].decode()
    parts['xl/worksheets/sheet1.xml'] = sheet.replace('</worksheet>', SPARKLINE + '</worksheet>').encode()
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return path


def sheet_xml(path):
    with zipfile.ZipFile(path) as archive:
        return archive.read('xl/worksheets/sheet1.xml').decode()


def test_preflight_flags_a_workbook_openpyxl_would_damage(tmp_path):
    workbook(tmp_path / 'risky.xlsx')
    workbook(tmp_path / 'plain.xlsx', sparkline=False)

    risky = run('preflight.py', 'risky.xlsx', cwd=tmp_path)
    plain = run('preflight.py', 'plain.xlsx', cwd=tmp_path)

    assert risky.returncode == 2 and 'sparklines: 1' in risky.stdout and 'UNSAFE' in risky.stdout
    assert plain.returncode == 0 and 'rich_text=True' in plain.stdout


def test_cells_edits_values_and_keeps_every_other_part(tmp_path):
    workbook(tmp_path / 'original.xlsx')

    result = run(
        'cells.py', 'original.xlsx', 'edited.xlsx',
        '--set', 'Budget!A4=Total', '--set', 'Budget!B4==SUM(B2:B3)', '--style', 'Budget!B4=Budget!B3',
        cwd=tmp_path,
    )

    assert result.returncode == 0, result.stdout
    assert '<x14:sparkline>' in sheet_xml(tmp_path / 'edited.xlsx')
    ws = load_workbook(tmp_path / 'edited.xlsx', rich_text=True)['Budget']
    assert ws['A4'].value == 'Total'
    assert ws['B4'].value == '=SUM(B2:B3)'
    assert ws['B4'].font.name == 'Arial'
    assert isinstance(ws['A5'].value, CellRichText)
    verdict = verify_structure.verify(
        tmp_path / 'original.xlsx', tmp_path / 'edited.xlsx',
        [verify_structure.parse_allow('Budget!A4:B4')],
    )
    assert verdict['status'] == 'success', verdict['issues']


def test_an_openpyxl_save_is_reported_with_what_it_lost(tmp_path):
    original = workbook(tmp_path / 'original.xlsx')
    wb = load_workbook(original)
    wb['Budget']['A4'] = 'Total'
    wb.save(tmp_path / 'edited.xlsx')

    verdict = verify_structure.verify(original, tmp_path / 'edited.xlsx', [verify_structure.parse_allow('Budget!A4')])

    details = [issue['detail'] for issue in verdict['issues'] if issue['kind'] == 'media']
    assert 'sparklines: 1 in the original, 0 in the output' in details
    assert 'rich-text cells: 1 in the original, 0 in the output' in details


@pytest.mark.parametrize('before, after', [
    ('_(* #,##0_);_(* (#,##0);_(* "-"_);_(@_)', r'_(* #,##0_);_(* \(#,##0\);_(* \-_);_(@_)'),
    ('0.0%;[Red]-0.0%', r'0.0%;[RED]\-0.0%'),
])
def test_libreoffice_spellings_of_a_number_format_mean_the_same(before, after):
    assert verify_structure.format_key(before) == verify_structure.format_key(after)


def test_a_different_number_format_still_differs():
    assert verify_structure.format_key('0.0%') != verify_structure.format_key('0%')


@pytest.mark.skipif(not shutil.which('soffice'), reason='needs LibreOffice')
def test_recalc_writes_values_without_rewriting_the_workbook(tmp_path):
    workbook(tmp_path / 'original.xlsx')
    run('cells.py', 'original.xlsx', 'edited.xlsx', '--set', 'Budget!B4==SUM(B2:B3)', cwd=tmp_path)

    result = run('recalc.py', tmp_path / 'edited.xlsx', cwd=tmp_path)

    assert result.returncode == 0, result.stdout + result.stderr
    assert '<x14:sparkline>' in sheet_xml(tmp_path / 'edited.xlsx')
    assert load_workbook(tmp_path / 'edited.xlsx', data_only=True)['Budget']['B4'].value == 255
    with zipfile.ZipFile(tmp_path / 'edited.xlsx') as archive:
        assert 'fullCalcOnLoad="1"' in archive.read('xl/workbook.xml').decode()
        assert 'Arial' in archive.read('xl/styles.xml').decode()
