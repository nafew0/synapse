"""Tests for skill/xlsx/scripts/verify_structure.py and copy_sheet.py.

Kept outside skill/xlsx/ on purpose: every file in a deployment skill directory is uploaded into
each user's code sandbox.

Run with openpyxl and pytest installed:
    python -m pytest tests/skill/xlsx
The LibreOffice round-trip test is skipped when `soffice` is not on PATH.
"""
import json
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

SCRIPTS = Path(__file__).resolve().parents[3] / 'skill' / 'xlsx' / 'scripts'
SCRIPT = SCRIPTS / 'verify_structure.py'
SHEET = 'Sheet 2'

sys.path.insert(0, str(SCRIPTS))
import copy_sheet  # noqa: E402
import verify_structure  # noqa: E402
from layout import layout_key  # noqa: E402


def make_form(path):
    """A requisition form shaped like the ones users send: merged header, bordered table, totals."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    # Excel writes borders in its automatic color, which LibreOffice re-saves as `auto`.
    thin = Side(style='thin', color=Color(indexed=64))
    ws['A2'] = 'Research and Education Network Trust'
    ws['A2'].font = Font(name='Arial', bold=True, size=14)
    ws.merge_cells('A2:F2')
    ws['A4'] = 'REQUISITION FORM'
    ws.merge_cells('A4:F4')
    for column, heading in zip('ABCDEF', ['Sl. #', 'Description', 'Quantity', 'Unit Cost', 'Estimated Cost', 'Budget Line']):
        ws[f'{column}9'] = heading
        ws[f'{column}9'].fill = PatternFill('solid', fgColor='00D9D9D9')
        ws[f'{column}9'].border = Border(top=thin, bottom=thin, left=thin, right=thin)
        ws.merge_cells(f'{column}9:{column}10')
    for row in (12, 13, 14):
        ws[f'A{row}'] = row - 11
        ws[f'E{row}'] = f'=D{row}*C{row}'
    ws['B12'] = 'USB hub'
    ws['C12'] = 1
    ws['D12'] = 4999
    ws['B18'] = 'Total Estimated cost'
    ws.merge_cells('B18:C18')
    ws['E18'] = '=SUM(E12:E17)'
    ws['A19'] = 'Amount in words: Four thousand nine hundred ninety-nine only'
    ws.merge_cells('A19:F19')
    ws.column_dimensions['B'].width = 42
    ws.column_dimensions['G'].hidden = True
    ws.row_dimensions[2].height = 28
    ws.oddHeader.center.text = 'Requisition Form'
    ws.oddFooter.center.text = 'Page &P of &N'
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.header = 0.25
    ws.print_title_rows = '9:10'
    ws.freeze_panes = 'A11'
    validation = DataValidation(type='whole', operator='greaterThan', formula1='0')
    ws.add_data_validation(validation)
    validation.add('C12:C14')
    wb.save(path)
    return path


@pytest.fixture
def original(tmp_path):
    return make_form(tmp_path / 'original.xlsx')


def edit(original, tmp_path, change):
    output = tmp_path / 'output.xlsx'
    shutil.copy(original, output)
    wb = load_workbook(output)
    change(wb[SHEET])
    wb.save(output)
    return output


def fill_items(ws):
    ws['B12'] = 'Monitor'
    ws['D12'] = 30000
    ws['B13'] = 'Laptop'
    ws['C13'] = 1
    ws['D13'] = 165000
    ws['A19'] = 'Amount in words: One lakh ninety-five thousand only'


ITEM_RANGES = [f'{SHEET}!B12:D13', f'{SHEET}!A19']


def run(original, output, allow, new_sheets=()):
    allowed = [verify_structure.parse_allow(spec) for spec in allow]
    declared = dict(verify_structure.parse_new_sheet(spec) for spec in new_sheets)
    return verify_structure.verify(original, output, allowed, declared)


COPY = 'Committee'
COPY_RANGES = [f'{COPY}!B12:D13', f'{COPY}!A19']


def add_copy(original, tmp_path, change=fill_items):
    """Add a tab copied from the form, the way a user asks for "a similar sheet alongside"."""
    output = tmp_path / 'output.xlsx'
    copy_sheet.copy_sheet(original, output, SHEET, COPY)
    wb = load_workbook(output)
    change(wb[COPY])
    wb.save(output)
    return output


def kinds(result):
    return {issue['kind'] for issue in result['issues']}


def test_an_edit_inside_the_declared_ranges_passes(original, tmp_path):
    output = edit(original, tmp_path, fill_items)

    result = run(original, output, ITEM_RANGES)

    assert result['status'] == 'success'
    assert result['cells_changed_inside_allowed'] == 6
    assert result['merged_ranges_checked'] == 10


def test_unmerging_the_form_fails(original, tmp_path):
    def unmerge_everything(ws):
        for merged in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(merged))
        fill_items(ws)

    result = run(original, edit(original, tmp_path, unmerge_everything), ITEM_RANGES)

    assert result['status'] == 'changes_found'
    removed = [issue['detail'] for issue in result['issues'] if issue['kind'] == 'unmerged']
    assert len(removed) == 10
    assert 'merged range A2:F2 was removed' in removed


def test_writing_into_the_header_fails_and_names_the_cell(original, tmp_path):
    def write_in_header(ws):
        ws['B6'] = 'Monitor'

    result = run(original, edit(original, tmp_path, write_in_header), ITEM_RANGES)

    assert result['status'] == 'changes_found'
    assert any(issue['kind'] == 'value' and issue['detail'].startswith('B6:') for issue in result['issues'])


def test_hardcoding_a_total_over_its_formula_fails(original, tmp_path):
    def hardcode_total(ws):
        fill_items(ws)
        ws['E18'] = 195000

    result = run(original, edit(original, tmp_path, hardcode_total), ITEM_RANGES)

    assert any(issue['detail'].startswith("E18: '=SUM(E12:E17)'") for issue in result['issues'])


def test_restyling_a_cell_outside_the_ranges_fails(original, tmp_path):
    def restyle_heading(ws):
        ws['B9'].font = Font(name='Calibri', size=9)

    result = run(original, edit(original, tmp_path, restyle_heading), ITEM_RANGES)

    assert kinds(result) == {'style'}


def test_formatting_inside_the_ranges_is_allowed(original, tmp_path):
    def restyle_item(ws):
        fill_items(ws)
        ws['D12'].number_format = '#,##0.00'

    assert run(original, edit(original, tmp_path, restyle_item), ITEM_RANGES)['status'] == 'success'


def test_resizing_a_column_fails(original, tmp_path):
    def widen(ws):
        ws.column_dimensions['B'].width = 60

    result = run(original, edit(original, tmp_path, widen), ITEM_RANGES)

    assert any(issue['kind'] == 'dimension' and 'column B' in issue['detail'] for issue in result['issues'])


def test_libreoffice_rounding_of_row_heights_is_tolerated(original, tmp_path):
    def round_like_libreoffice(ws):
        ws.row_dimensions[2].height = 27.75
        ws.row_dimensions[12].height = 15.0

    assert run(original, edit(original, tmp_path, round_like_libreoffice), [])['status'] == 'success'


def test_an_alpha_byte_on_a_fill_color_is_tolerated(original, tmp_path):
    def set_alpha(ws):
        ws['B9'].fill = PatternFill('solid', fgColor='FFD9D9D9')

    assert run(original, edit(original, tmp_path, set_alpha), [])['status'] == 'success'


def test_a_different_fill_color_fails(original, tmp_path):
    def recolor(ws):
        ws['B9'].fill = PatternFill('solid', fgColor='FFFF0000')

    assert kinds(run(original, edit(original, tmp_path, recolor), [])) == {'style'}


def test_explicit_default_alignment_is_tolerated(original, tmp_path):
    def spell_out_defaults(ws):
        ws['A12'].alignment = Alignment(horizontal='general', vertical='bottom')

    assert run(original, edit(original, tmp_path, spell_out_defaults), [])['status'] == 'success'


def test_removing_data_validation_fails(original, tmp_path):
    def drop_validation(ws):
        ws.data_validations.dataValidation = []

    assert 'data_validation' in kinds(run(original, edit(original, tmp_path, drop_validation), []))


def test_an_unused_second_bound_on_validation_is_tolerated(original, tmp_path):
    def add_unused_bound(ws):
        ws.data_validations.dataValidation[0].formula2 = '0'

    assert run(original, edit(original, tmp_path, add_unused_bound), [])['status'] == 'success'


def test_renaming_a_sheet_fails(original, tmp_path):
    def rename(ws):
        ws.title = 'Requisition'

    assert 'sheets' in kinds(run(original, edit(original, tmp_path, rename), []))


def test_the_automatic_border_color_matches_no_color(original, tmp_path):
    def drop_automatic_color(ws):
        ws['A9'].border = Border(top=Side(style='thin'), bottom=Side(style='thin'), left=Side(style='thin'), right=Side(style='thin'))

    assert run(original, edit(original, tmp_path, drop_automatic_color), [])['status'] == 'success'


def test_a_different_border_color_fails(original, tmp_path):
    def recolor_border(ws):
        red = Side(style='thin', color='FFFF0000')
        ws['A9'].border = Border(top=red, bottom=red, left=red, right=red)

    assert kinds(run(original, edit(original, tmp_path, recolor_border), [])) == {'style'}


def test_changing_the_page_setup_fails(original, tmp_path):
    wb = load_workbook(original)
    wb[SHEET].page_setup.orientation = 'portrait'
    wb[SHEET].sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(original)

    def turn_landscape(ws):
        ws.page_setup.orientation = 'landscape'
        ws.sheet_properties.pageSetUpPr.fitToPage = False

    result = run(original, edit(original, tmp_path, turn_landscape), [])

    assert [issue['detail'] for issue in result['issues']] == [
        "orientation: 'portrait' -> 'landscape'",
        'fit to page: True -> False',
    ]


def test_an_undeclared_new_sheet_fails(original, tmp_path):
    assert kinds(run(original, add_copy(original, tmp_path), [])) == {'sheets'}


def test_a_bare_copy_worksheet_fails_on_everything_it_leaves_out(original, tmp_path):
    output = tmp_path / 'output.xlsx'
    wb = load_workbook(original)
    wb.copy_worksheet(wb[SHEET]).title = COPY
    wb.save(output)

    result = run(original, output, [], [f'{COPY}:{SHEET}'])

    assert {(issue['kind'], issue['sheet']) for issue in result['issues']} == {('data_validation', COPY), ('layout', COPY)}
    lost = {issue['detail'].split(':')[0] for issue in result['issues'] if issue['kind'] == 'layout'}
    assert {'print titles', 'header and footer', 'freeze panes'} <= lost


def test_setting_an_unset_page_option_is_tolerated(original, tmp_path):
    def choose_like_libreoffice(ws):
        ws.page_setup.orientation = 'portrait'
        ws.page_setup.paperSize = 9

    assert run(original, edit(original, tmp_path, choose_like_libreoffice), [])['status'] == 'success'


def test_a_copied_tab_filled_inside_its_ranges_passes(original, tmp_path):
    result = run(original, add_copy(original, tmp_path), COPY_RANGES, [f'{COPY}:{SHEET}'])

    assert result['status'] == 'success'
    assert result['sheets_checked'] == 2
    assert result['merged_ranges_checked'] == 20
    assert result['cells_changed_inside_allowed'] == 6


def test_a_copied_tab_that_departs_from_its_template_fails(original, tmp_path):
    def rebuild_loosely(ws):
        fill_items(ws)
        ws.unmerge_cells('A19:F19')
        ws['B18'].font = Font(name='Arial', size=10)

    result = run(original, add_copy(original, tmp_path, rebuild_loosely), COPY_RANGES, [f'{COPY}:{SHEET}'])

    assert {(issue['kind'], issue['sheet']) for issue in result['issues']} == {('unmerged', COPY), ('style', COPY)}


def test_inserting_rows_into_a_copied_tab_fails(original, tmp_path):
    """openpyxl's insert_rows moves values but not merged ranges, formulas or row heights."""

    def insert_rows(ws):
        ws.insert_rows(15, 3)

    result = run(original, add_copy(original, tmp_path, insert_rows), [f'{COPY}!12:17'], [f'{COPY}:{SHEET}'])

    assert any(issue['detail'].startswith("B18: 'Total Estimated cost' -> None") for issue in result['issues'])


def test_a_new_sheet_without_a_template_only_has_to_exist(original, tmp_path):
    def add_notes(ws):
        ws.parent.create_sheet('Notes')['A1'] = 'Prepared from the committee memo'

    output = edit(original, tmp_path, add_notes)

    assert run(original, output, [], ['Notes'])['status'] == 'success'
    assert run(original, original, [], ['Missing'])['issues'][0]['detail'].startswith('declared new sheet')


def test_new_sheets_may_go_anywhere_around_the_original_sheets(original, tmp_path):
    def add_around(ws):
        ws.parent.create_sheet('Notes', 0)
        ws.parent.create_sheet('Summary')

    assert run(original, edit(original, tmp_path, add_around), [], ['Notes', 'Summary'])['status'] == 'success'


def test_reordering_the_original_sheets_fails(original, tmp_path):
    def add_and_reorder(ws):
        ws.parent.create_sheet('Budget')
        ws.parent.move_sheet(SHEET, offset=1)

    assert kinds(run(original, edit(original, tmp_path, add_and_reorder), [])) == {'sheets'}


@pytest.mark.parametrize(
    ('declared', 'message'),
    [([SHEET], 'already exist'), ([f'{COPY}:Sheet 9'], 'Sheet 9')],
)
def test_bad_new_sheet_declarations_are_rejected(original, tmp_path, declared, message):
    with pytest.raises(verify_structure.InputError, match=message):
        run(original, add_copy(original, tmp_path), [], declared)


@pytest.mark.parametrize(('spec', 'expected'), [('Notes', ('Notes', None)), ("'Committee':'Sheet 2'", (COPY, SHEET))])
def test_new_sheet_specs_accept_a_name_and_an_optional_template(spec, expected):
    assert verify_structure.parse_new_sheet(spec) == expected


@pytest.mark.parametrize('spec', ['', ':Sheet 2', 'Committee:'])
def test_malformed_new_sheet_specs_are_rejected(spec):
    with pytest.raises(verify_structure.InputError):
        verify_structure.parse_new_sheet(spec)


def test_losing_an_image_fails(original, tmp_path):
    with_logo = tmp_path / 'with_logo.xlsx'
    shutil.copy(original, with_logo)
    with zipfile.ZipFile(with_logo, 'a') as archive:
        archive.writestr('xl/media/image1.png', b'\x89PNG logo')

    result = run(with_logo, original, [])

    assert result['status'] == 'changes_found'
    assert any(issue['kind'] == 'media' and issue['detail'].startswith('images: 1') for issue in result['issues'])


@pytest.mark.parametrize(
    ('spec', 'expected'),
    [
        ("Sheet 2!B12", ('Sheet 2', 2, 12, 2, 12)),
        ("'Sheet 2'!D13:B12", ('Sheet 2', 2, 12, 4, 13)),
        ('Sheet 2!$B$12:$D$14', ('Sheet 2', 2, 12, 4, 14)),
        ('Sheet 2!12:14', ('Sheet 2', 1, 12, verify_structure.UNBOUNDED, 14)),
        ('Sheet 2!b:d', ('Sheet 2', 2, 1, 4, verify_structure.UNBOUNDED)),
    ],
)
def test_allow_ranges_accept_cells_blocks_rows_and_columns(spec, expected):
    assert verify_structure.parse_allow(spec) == expected


def test_a_whole_row_allowance_covers_every_column(original, tmp_path):
    def fill_whole_row(ws):
        ws['F12'] = '12000500'
        ws['B12'] = 'Monitor'

    assert run(original, edit(original, tmp_path, fill_whole_row), [f'{SHEET}!12:12'])['status'] == 'success'


@pytest.mark.parametrize('spec', ['B12:D14', 'Sheet 2!', 'Sheet 2!not-a-range'])
def test_malformed_allow_ranges_are_rejected(spec):
    with pytest.raises(verify_structure.InputError):
        verify_structure.parse_allow(spec)


def test_an_allowance_for_an_unknown_sheet_is_rejected(original, tmp_path):
    output = edit(original, tmp_path, fill_items)

    with pytest.raises(verify_structure.InputError, match='Sheet 3'):
        run(original, output, ['Sheet 3!A1'])


def change_margin(ws):
    ws.page_margins.left = 1.0


def drop_header(ws):
    ws.oddHeader.center.text = None


def unhide_column(ws):
    ws.column_dimensions['G'].hidden = False


def unfreeze(ws):
    ws.freeze_panes = None


def move_print_titles(ws):
    ws.print_title_rows = '9:9'


@pytest.mark.parametrize(
    ('change', 'setting'),
    [
        (change_margin, 'left margin'),
        (drop_header, 'header and footer'),
        (unhide_column, 'hidden columns'),
        (unfreeze, 'freeze panes'),
        (move_print_titles, 'print titles'),
    ],
)
def test_changing_the_print_layout_fails(original, tmp_path, change, setting):
    result = run(original, edit(original, tmp_path, change), [])

    assert kinds(result) == {'layout'}
    assert any(issue['detail'].startswith(f'{setting}:') for issue in result['issues'])


def test_copy_sheet_keeps_what_copy_worksheet_leaves_out(original, tmp_path):
    wb = load_workbook(original)
    wb.create_sheet('Notes')
    wb.save(original)
    output = tmp_path / 'output.xlsx'

    result = copy_sheet.copy_sheet(original, output, SHEET, COPY)

    copied = load_workbook(output)
    assert copied.sheetnames == [SHEET, COPY, 'Notes']
    assert layout_key(copied[COPY]) == layout_key(copied[SHEET])
    assert set(result['also_copied']) == {'header/footer', 'print title rows', 'data validation'}
    assert result['verify_with'] == f"--new-sheet '{COPY}:{SHEET}'"


@pytest.mark.parametrize(('name', 'message'), [(SHEET, 'already exists'), ('Q1/Q2', 'Invalid'), ('x' * 32, 'Invalid')])
def test_copy_sheet_rejects_unusable_names(original, tmp_path, name, message):
    with pytest.raises(copy_sheet.InputError, match=message):
        copy_sheet.copy_sheet(original, tmp_path / 'output.xlsx', SHEET, name)


def test_copy_sheet_cli_refuses_to_overwrite_the_original(original):
    completed = subprocess.run(
        [sys.executable, str(SCRIPTS / 'copy_sheet.py'), str(original), str(original), '--template', SHEET, '--name', COPY],
        capture_output=True,
        text=True,
    )

    assert completed.returncode == 1
    assert 'same file' in json.loads(completed.stdout)['error']


def fit_issues(result):
    return {issue['detail'].split(':')[0]: issue['detail'] for issue in result['issues'] if issue['kind'] == 'fit'}


def test_a_number_wider_than_its_column_fails_as_hashes(original, tmp_path):
    def big_number(ws):
        ws['D12'] = 123456789012.5
        ws['D12'].number_format = '#,##0.00'

    result = run(original, edit(original, tmp_path, big_number), ITEM_RANGES)

    assert list(fit_issues(result)) == ['D12']
    assert fit_issues(result)['D12'].endswith('prints as ####')


def test_text_longer_than_a_merged_cell_fails(original, tmp_path):
    def long_words(ws):
        ws['A19'] = 'Amount in words: ' + 'one hundred and ninety-five thousand taka only, ' * 3

    assert list(fit_issues(run(original, edit(original, tmp_path, long_words), ITEM_RANGES))) == ['A19']


def test_long_text_fails_only_when_its_neighbour_blocks_it(original, tmp_path):
    description = 'Honorarium of the investigation committee for the second meeting'

    def blocked(ws):
        ws['B12'] = description

    def spilling(ws):
        ws['B12'] = description
        ws['C12'] = None

    ranges = [f'{SHEET}!B12:C12']
    assert list(fit_issues(run(original, edit(original, tmp_path, blocked), ranges))) == ['B12']
    assert run(original, edit(original, tmp_path, spilling), ranges)['status'] == 'success'


def test_wrapped_text_fails_only_in_a_row_of_fixed_height(original, tmp_path):
    def wrapped(fixed):
        def change(ws):
            ws['B14'] = 'Honorarium of the investigation committee members for four meetings, less AIT'
            ws['B14'].alignment = Alignment(wrap_text=True)
            if fixed:
                ws.row_dimensions[14].height = 15

        return change

    ranges = [f'{SHEET}!B14', f'{SHEET}!14:14']
    assert run(original, edit(original, tmp_path, wrapped(False)), ranges)['status'] == 'success'
    assert list(fit_issues(run(original, edit(original, tmp_path, wrapped(True)), ranges))) == ['B14']


def test_overflow_the_original_already_had_is_not_reported(original, tmp_path):
    wb = load_workbook(original)
    wb[SHEET]['A4'] = 'REQUISITION FORM FOR THE PROCUREMENT OF GOODS AND SERVICES UNDER THE ANNUAL BUDGET ' * 2
    wb.save(original)

    assert run(original, edit(original, tmp_path, fill_items), ITEM_RANGES)['status'] == 'success'


def test_content_on_a_new_sheet_must_fit_too(original, tmp_path):
    def add_summary(ws):
        summary = ws.parent.create_sheet('Summary')
        summary['A1'] = 1234567.891
        summary['A1'].number_format = '#,##0.00'
        summary['B1'] = 'next'

    result = run(original, edit(original, tmp_path, add_summary), [], ['Summary'])

    assert [(issue['kind'], issue['sheet']) for issue in result['issues']] == [('fit', 'Summary')]


def cli(*args):
    completed = subprocess.run([sys.executable, str(SCRIPT), *map(str, args)], capture_output=True, text=True)
    return completed.returncode, json.loads(completed.stdout)


def test_cli_exits_zero_on_success(original, tmp_path):
    code, payload = cli(original, edit(original, tmp_path, fill_items), '--allow', ITEM_RANGES[0], '--allow', ITEM_RANGES[1])

    assert (code, payload['status']) == (0, 'success')


def test_cli_exits_two_when_changes_are_found(original, tmp_path):
    code, payload = cli(original, edit(original, tmp_path, fill_items))

    assert (code, payload['status']) == (2, 'changes_found')


def test_cli_refuses_to_compare_a_file_with_itself(original):
    code, payload = cli(original, original)

    assert code == 1
    assert 'same file' in payload['error']


@pytest.mark.skipif(shutil.which('soffice') is None, reason='LibreOffice is not installed')
def test_a_correct_edit_still_passes_after_libreoffice_recalculates_it(original, tmp_path):
    output = edit(original, tmp_path, fill_items)
    recalc = subprocess.run(
        [sys.executable, str(SCRIPTS / 'recalc.py'), str(output), '90'], capture_output=True, text=True
    )
    assert json.loads(recalc.stdout).get('status') == 'success', recalc.stdout + recalc.stderr

    assert run(original, output, ITEM_RANGES)['status'] == 'success'


@pytest.mark.skipif(shutil.which('soffice') is None, reason='LibreOffice is not installed')
def test_a_copied_tab_still_passes_after_libreoffice_recalculates_it(original, tmp_path):
    output = add_copy(original, tmp_path)
    recalc = subprocess.run(
        [sys.executable, str(SCRIPTS / 'recalc.py'), str(output), '90'], capture_output=True, text=True
    )
    assert json.loads(recalc.stdout).get('status') == 'success', recalc.stdout + recalc.stderr

    assert run(original, output, COPY_RANGES, [f'{COPY}:{SHEET}'])['status'] == 'success'
