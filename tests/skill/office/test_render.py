"""Tests for office/render.py, gate step 5: render with LibreOffice, read back with Tesseract.

The comparison tests are pure Python. The render tests need LibreOffice, Poppler and Tesseract
with its Bengali model plus Nikosh, and skip without them. The Bijoy table test renders the
hand-typed Bijoy strings in an ANSI (Bijoy-layout) font and reads them back, so it checks the
table against real glyphs; it needs Kalpurush ANSI or Siyam Rupali ANSI.

Real files are not committed: they are users' uploads. Point BANGLA_CORPUS at a directory holding
the files named in `corpus.json` to run them against their recorded results:
    BANGLA_CORPUS=/path/to/files python -m pytest tests/skill/office/test_render.py
"""
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

import pytest

from test_bijoy import PAIRS
from test_fix_verify import TestFixPptx, bijoy_docx, letter_docx
from test_policy import HAS_NIKOSH, SCRIPTS, SENTENCE

sys.path.insert(0, str(SCRIPTS))
from office import bangla, bijoy, render  # noqa: E402
from office.fix_bangla import fix  # noqa: E402
from office.verify_bangla import verify  # noqa: E402

docx = pytest.importorskip('docx')
openpyxl = pytest.importorskip('openpyxl')
from docx.shared import Pt  # noqa: E402

CAN_RENDER = HAS_NIKOSH and not render.missing_tools()
needs_render = pytest.mark.skipif(not CAN_RENDER, reason='needs LibreOffice, Poppler, Tesseract ben and Nikosh')
ANSI_FONT = next((name for name in ('Siyam Rupali ANSI', 'Kalpurush ANSI') if bangla.font_path(name)), None)
MODERN_ONLY = set('ÿøÍ')
"""SutonnyMJ codes the older ANSI fonts draw differently: ক্ষ ্ল ্ত are `¶ ¬ —` in them."""
CORPUS = Path(os.environ.get('BANGLA_CORPUS', '')) if os.environ.get('BANGLA_CORPUS') else None
EXPECTED = json.loads((Path(__file__).with_name('corpus.json')).read_text(encoding='utf-8'))


class TestCompare:
    def test_the_same_text_misses_nothing(self):
        assert render.compare(SENTENCE, SENTENCE)[0] == 0

    def test_word_order_does_not_matter(self):
        words = SENTENCE.split()
        assert render.compare(SENTENCE, ' '.join(reversed(words)))[0] == 0

    def test_a_repeated_word_must_be_read_every_time(self):
        rate, _, _ = render.compare('রেজিস্ট্রার ' * 10, 'রেজিস্ট্রার')
        assert rate == pytest.approx(0.9)

    def test_drawn_order_reads_as_a_miss(self):
        rate, _, unread = render.compare('বাংলাদেশ বিশ্ববিদ্যালয়', 'বাংলােদশ িবশ্বিবদ্যালয়')
        assert rate > 0.3
        assert unread == ['বাংলাদেশ', 'বিশ্ববিদ্যালয়']

    def test_inherited_garble_is_not_judged(self):
        garbled = 'বাংলােদশ'
        assert render.compare(f'{garbled} কমিশন', 'কমিশন', frozenset({bangla.normalize(garbled)}))[0] == 0

    def test_no_bangla_is_nothing_to_judge(self):
        assert render.compare('Memo No. 37', '') == (0.0, 0, [])


@needs_render
class TestRender:
    def test_a_fixed_new_document_reads_back(self, tmp_path):
        out = tmp_path / 'out.docx'
        fix(letter_docx(tmp_path / 'in.docx'), out, new=True)
        result = render.check(out)
        assert result['status'] == 'clean', result
        assert result['pages'][0]['miss_rate'] < 0.03

    def test_drawn_order_text_fails(self, tmp_path):
        path = tmp_path / 'broken.docx'
        document = docx.Document()
        for _ in range(4):
            run = document.add_paragraph().add_run('বাংলােদশ িবশ্বিবদ্যালয় মঞ্জুির কিমশন েসেপ্টম্বর')
            run.font.name = 'Nikosh'
            run.font.size = Pt(14)
        document.save(path)
        result = render.check(path)
        assert result['status'] == 'defects_found', result

    def test_the_same_garble_in_the_original_is_not_blamed(self, tmp_path):
        path = tmp_path / 'broken.docx'
        document = docx.Document()
        for _ in range(4):
            document.add_paragraph().add_run('বাংলােদশ িবশ্বিবদ্যালয় কিমশন').font.name = 'Nikosh'
        document.add_paragraph().add_run(SENTENCE).font.name = 'Nikosh'
        document.save(path)
        assert render.check(path, original=path)['status'] == 'clean'

    def test_a_clipped_cell_fails(self, tmp_path):
        path = tmp_path / 'narrow.xlsx'
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in range(1, 8):
            sheet.cell(row, 1, SENTENCE).font = openpyxl.styles.Font(name='Nikosh', size=12)
            sheet.cell(row, 2, 'x')
        sheet.column_dimensions['A'].width = 12
        workbook.save(path)
        assert render.check(path)['status'] == 'defects_found'
        sheet.column_dimensions['A'].width = 70
        workbook.save(path)
        assert render.check(path)['status'] == 'clean'

    def test_a_deck_reads_back(self, tmp_path):
        deck = TestFixPptx().deck(tmp_path / 'in.pptx', text=f'{SENTENCE} {SENTENCE}', font='Nikosh')
        assert render.check(deck)['status'] == 'clean'

    def test_a_converted_bijoy_letter_reads_back(self, tmp_path):
        out = tmp_path / 'out.docx'
        fix(bijoy_docx(tmp_path / 'in.docx'), out)
        result = verify(out, tmp_path / 'in.docx', render=True)
        assert result['status'] == 'clean', result
        assert result['render']['pages']

    def test_the_gate_runs_the_render_step(self, tmp_path):
        out = tmp_path / 'out.docx'
        fix(letter_docx(tmp_path / 'in.docx'), out, new=True)
        ran = subprocess.run(
            [sys.executable, str(SCRIPTS / 'office' / 'verify_bangla.py'), str(out)], capture_output=True, text=True
        )
        assert ran.returncode == 0, ran.stdout + ran.stderr
        assert json.loads(ran.stdout)['render']['status'] == 'clean'
        ran = subprocess.run(
            [sys.executable, str(SCRIPTS / 'office' / 'verify_bangla.py'), str(out), '--no-render'],
            capture_output=True, text=True,
        )
        assert 'render' not in json.loads(ran.stdout)


@needs_render
@pytest.mark.skipif(ANSI_FONT is None, reason='needs an ANSI (Bijoy layout) font: Kalpurush ANSI or Siyam Rupali ANSI')
def test_bijoy_table_matches_what_a_bijoy_font_draws(tmp_path):
    """Draw the Bijoy strings in a Bijoy-layout font and read them back: the pixels must say what
    `to_unicode` says the codes mean."""
    typed = [text for text, _ in PAIRS if not MODERN_ONLY & set(text)]
    path = tmp_path / 'bijoy.docx'
    document = docx.Document()
    for text in typed:
        run = document.add_paragraph().add_run(text)
        run.font.name = ANSI_FONT
        run.font.size = Pt(16)
    document.save(path)
    with tempfile.TemporaryDirectory() as scratch:
        pdf = render.to_pdf(path, Path(scratch))
        seen = render.read_page(pdf, 1, Path(scratch))
    rate, _, unread = render.compare('\n'.join(bijoy.to_unicode(text) for text in typed), seen)
    assert rate < render.MISS_THRESHOLD, (rate, unread)


@needs_render
@pytest.mark.skipif(CORPUS is None, reason='set BANGLA_CORPUS to a directory of the real files in corpus.json')
@pytest.mark.parametrize('name', sorted(EXPECTED))
def test_real_files(name):
    path = CORPUS / name
    if not path.is_file():
        pytest.skip(f'{name} is not in BANGLA_CORPUS')
    expected = EXPECTED[name]
    result = render.check(path, max_pages=20)
    assert result['status'] == expected['render'], result
    if path.suffix == '.pdf':
        return
    assert verify(path)['status'] == expected['verify']


@needs_render
@pytest.mark.skipif(CORPUS is None, reason='set BANGLA_CORPUS to a directory of the real files in corpus.json')
@pytest.mark.parametrize('name', sorted(name for name in EXPECTED if not name.endswith('.pdf')))
def test_real_files_after_the_fix(name, tmp_path):
    """The whole pipeline a specialist runs: fix_bangla.py, then every gate step against the
    original. Faults the original had are notes; only the edit's own faults fail."""
    path = CORPUS / name
    if not path.is_file():
        pytest.skip(f'{name} is not in BANGLA_CORPUS')
    out = tmp_path / name
    fix(path, out)
    result = verify(out, path, render=True, max_pages=20)
    assert result['status'] == EXPECTED[name]['fixed'], result
