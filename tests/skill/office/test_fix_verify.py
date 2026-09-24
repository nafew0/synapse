"""Tests for fix_bangla.py and verify_bangla.py on real DOCX, PPTX and XLSX files.

Files are built with python-docx, python-pptx and openpyxl, then fixed and checked through the
scripts' own functions and command lines. Font-dependent tests skip without the sandbox fonts.
"""
import io
import json
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from lxml import etree

from test_policy import HAS_NIKOSH, SCRIPTS, SENTENCE, needs_fonts

sys.path.insert(0, str(SCRIPTS))
from office import bangla  # noqa: E402
from office.fix_bangla import fix, obfuscate  # noqa: E402
from office.runs import NS, W  # noqa: E402
from office.verify_bangla import verify  # noqa: E402

docx = pytest.importorskip('docx')
pptx = pytest.importorskip('pptx')
openpyxl = pytest.importorskip('openpyxl')
from docx.shared import Pt  # noqa: E402
from pptx.util import Inches  # noqa: E402

OFFICE = SCRIPTS / 'office'
LETTER = 'বিষয়: রোকেয়া চেয়ার ২০২৬ নিয়োগের নিমিত্ত মনোনয়ন আহ্বান।'


def part(path, name):
    with zipfile.ZipFile(path) as archive:
        return archive.read(name)


def xml(path, name):
    return etree.fromstring(part(path, name))


def letter_docx(path, font='Calibri'):
    document = docx.Document()
    heading = document.add_heading('', level=1)
    heading.add_run(LETTER)
    body = document.add_paragraph()
    run = body.add_run(SENTENCE)
    run.font.name = font
    run.font.size = Pt(13)
    run.bold = True
    document.add_paragraph('Memo No. 37.01.0000 — English stays in its own font.')
    document.save(path)
    return path


def bangla_runs(path):
    root = xml(path, 'word/document.xml')
    return [run for run in root.iter(f'{{{W}}}r') if bangla.has_bangla(''.join(run.itertext()))]


class TestFixDocx:
    def test_bangla_runs_get_the_complex_script_slots(self, tmp_path):
        out = tmp_path / 'out.docx'
        fix(letter_docx(tmp_path / 'in.docx'), out, new=True, embed=False)
        body = bangla_runs(out)[1]
        rpr = body.find('w:rPr', NS)
        fonts = rpr.find('w:rFonts', NS)
        assert fonts.get(f'{{{W}}}cs') == 'Nikosh'
        assert fonts.get(f'{{{W}}}ascii') == 'Calibri'
        assert rpr.find('w:szCs', NS).get(f'{{{W}}}val') == '26'
        assert rpr.find('w:bCs', NS) is not None
        assert rpr.find('w:lang', NS).get(f'{{{W}}}bidi') == 'bn-BD'
        order = [etree.QName(child).localname for child in rpr]
        assert order == sorted(order, key=['rFonts', 'b', 'bCs', 'sz', 'szCs', 'lang'].index)

    def test_english_runs_are_untouched(self, tmp_path):
        source = letter_docx(tmp_path / 'in.docx')
        out = tmp_path / 'out.docx'
        fix(source, out, new=True, embed=False)
        english = [r for r in xml(out, 'word/document.xml').iter(f'{{{W}}}r') if 'Memo No.' in ''.join(r.itertext())]
        assert english[0].find('w:rPr', NS) is None

    def test_styles_get_complex_script_sizes(self, tmp_path):
        out = tmp_path / 'out.docx'
        fix(letter_docx(tmp_path / 'in.docx'), out, new=True, embed=False)
        styles = xml(out, 'word/styles.xml')
        for rpr in styles.iter(f'{{{W}}}rPr'):
            size = rpr.find('w:sz', NS)
            if size is not None:
                assert rpr.find('w:szCs', NS).get(f'{{{W}}}val') == size.get(f'{{{W}}}val')

    def test_an_edit_keeps_an_inherited_font_that_draws_bangla(self, tmp_path):
        source = letter_docx(tmp_path / 'in.docx')
        heading_font = bangla_runs(source)[0].find('w:rPr/w:rFonts', NS)
        out = tmp_path / 'out.docx'
        fix(source, out, new=False, embed=False)
        fonts = bangla_runs(out)[0].find('w:rPr/w:rFonts', NS)
        assert (fonts.get(f'{{{W}}}cs') if fonts is not None else None) == (
            heading_font.get(f'{{{W}}}cs') if heading_font is not None else None
        )

    def test_a_bijoy_font_is_replaced_even_in_an_edit(self, tmp_path):
        source = tmp_path / 'in.docx'
        document = docx.Document()
        run = document.add_paragraph().add_run(SENTENCE)
        run._r.get_or_add_rPr().get_or_add_rFonts().set(f'{{{W}}}cs', 'SutonnyMJ')
        document.save(source)
        out = tmp_path / 'out.docx'
        fix(source, out, new=False, embed=False)
        assert bangla_runs(out)[0].find('w:rPr/w:rFonts', NS).get(f'{{{W}}}cs') == 'Nikosh'

    def test_the_result_opens(self, tmp_path):
        out = tmp_path / 'out.docx'
        fix(letter_docx(tmp_path / 'in.docx'), out, new=True, embed=HAS_NIKOSH)
        assert bangla.compose(SENTENCE) in [p.text for p in docx.Document(out).paragraphs]


@needs_fonts
class TestEmbedding:
    def test_nikosh_is_embedded_whole_and_decodes(self, tmp_path):
        out = tmp_path / 'out.docx'
        result = fix(letter_docx(tmp_path / 'in.docx'), out, new=True)
        assert result['embedded'] is True
        table = xml(out, 'word/fontTable.xml')
        embed = table.find("w:font[@w:name='Nikosh']/w:embedRegular", NS)
        key = embed.get(f'{{{W}}}fontKey')
        rels = xml(out, 'word/_rels/fontTable.xml.rels')
        target = next(r.get('Target') for r in rels if r.get('Id') == embed.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'))
        stored = part(out, f'word/{target}')
        assert obfuscate(stored, key) == Path(bangla.font_path('Nikosh')).read_bytes()
        assert xml(out, 'word/settings.xml').find('w:embedTrueTypeFonts', NS) is not None
        assert b'Extension="odttf"' in part(out, '[Content_Types].xml')

    def test_obfuscation_is_its_own_inverse(self):
        key = '{7942CD47-D1DE-4D94-9644-A818D45CB7BA}'
        data = bytes(range(64))
        assert obfuscate(obfuscate(data, key), key) == data
        assert obfuscate(data, key)[:32] != data[:32]
        assert obfuscate(data, key)[32:] == data[32:]

    def test_fixing_twice_embeds_once(self, tmp_path):
        once, twice = tmp_path / 'once.docx', tmp_path / 'twice.docx'
        fix(letter_docx(tmp_path / 'in.docx'), once, new=True)
        fix(once, twice, new=True)
        with zipfile.ZipFile(twice) as archive:
            assert [n for n in archive.namelist() if n.endswith('.odttf')] == ['word/fonts/font1.odttf']

    def test_the_schema_validator_accepts_the_result(self, tmp_path):
        source = letter_docx(tmp_path / 'in.docx')
        out = tmp_path / 'out.docx'
        fix(source, out, new=True)
        checked = subprocess.run(
            [sys.executable, str(OFFICE / 'validate.py'), str(out), '--original', str(source)],
            capture_output=True, text=True,
        )
        assert checked.returncode == 0, checked.stdout + checked.stderr


class TestFixPptx:
    def deck(self, path, text=SENTENCE, font='Calibri'):
        presentation = pptx.Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
        box.text_frame.text = text
        box.text_frame.paragraphs[0].runs[0].font.name = font
        presentation.save(path)
        return path

    def test_bangla_runs_get_a_cs_and_language(self, tmp_path):
        out = tmp_path / 'out.pptx'
        fix(self.deck(tmp_path / 'in.pptx'), out, new=True)
        root = xml(out, 'ppt/slides/slide1.xml')
        rpr = root.find('.//a:r/a:rPr', NS)
        assert rpr.get('lang') == 'bn-BD'
        children = [etree.QName(child).localname for child in rpr]
        assert children.index('latin') < children.index('cs')
        assert rpr.find('a:cs', NS).get('typeface') == 'Nikosh'
        assert rpr.find('a:latin', NS).get('typeface') == 'Calibri'

    def test_the_result_opens(self, tmp_path):
        out = tmp_path / 'out.pptx'
        fix(self.deck(tmp_path / 'in.pptx'), out, new=True)
        assert pptx.Presentation(out).slides[0].shapes[0].text_frame.text == bangla.compose(SENTENCE)


class TestVerify:
    def test_a_fixed_new_document_is_clean(self, tmp_path):
        out = tmp_path / 'out.docx'
        fix(letter_docx(tmp_path / 'in.docx'), out, new=True, embed=HAS_NIKOSH)
        result = verify(out)
        if HAS_NIKOSH:
            assert result['status'] == 'clean', result
        else:
            assert {f['check'] for f in result['findings']} <= {'font'}

    def test_a_mismatched_complex_script_size_is_a_defect(self, tmp_path):
        result = verify(letter_docx(tmp_path / 'in.docx'))
        assert any('its w:szCs is not' in f['detail'] for f in result['findings'])

    def test_a_lost_paragraph_is_a_defect(self, tmp_path):
        source = letter_docx(tmp_path / 'in.docx')
        edited = tmp_path / 'edited.docx'
        document = docx.Document(source)
        paragraph = document.paragraphs[1]
        paragraph._p.getparent().remove(paragraph._p)
        document.save(edited)
        result = verify(edited, source)
        assert [f['check'] for f in result['findings'] if f['check'] == 'lost'] == ['lost']
        assert verify(edited, source, allowed=['কার্যালয়ের'])['findings'] == [
            f for f in result['findings'] if f['check'] != 'lost'
        ]

    def test_broken_new_text_is_a_defect(self, tmp_path):
        path = tmp_path / 'broken.docx'
        document = docx.Document()
        run = document.add_paragraph().add_run('বাংলােদশ িবদালয়')
        run._r.get_or_add_rPr().get_or_add_rFonts().set(f'{{{W}}}cs', 'Nikosh')
        document.save(path)
        assert 'garbled' in {f['check'] for f in verify(path)['findings']}

    def test_a_latin_font_in_a_new_cell_fails_but_old_cells_are_only_noted(self, tmp_path):
        source, edited = tmp_path / 'in.xlsx', tmp_path / 'out.xlsx'
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet['A1'] = 'তদন্ত কমিটি'
        sheet['A1'].font = openpyxl.styles.Font(name='Arial')
        workbook.save(source)
        sheet['A2'] = 'সম্মানী বিবরণী'
        sheet['A2'].font = openpyxl.styles.Font(name='Arial')
        sheet['A3'] = 'মোট'
        sheet['A3'].font = openpyxl.styles.Font(name='Nikosh')
        workbook.save(edited)
        result = verify(edited, source)
        assert [f['where'] for f in result['findings']] == ['Sheet!A2']
        assert [f['where'] for f in result['notes']] == ['Sheet!A1']

    def test_a_text_box_turned_into_a_picture_is_a_defect(self, tmp_path):
        source = TestFixPptx().deck(tmp_path / 'in.pptx', font='Nikosh')
        edited = tmp_path / 'out.pptx'
        presentation = pptx.Presentation(source)
        slide = presentation.slides[0]
        box = slide.shapes[0]
        box._element.getparent().remove(box._element)
        png = io.BytesIO()
        __import__('PIL.Image', fromlist=['Image']).new('RGB', (4, 4)).save(png, 'PNG')
        png.seek(0)
        slide.shapes.add_picture(png, Inches(1), Inches(1))
        presentation.save(edited)
        checks = {f['check'] for f in verify(edited, source)['findings']}
        assert {'lost', 'text'} <= checks

    def test_command_line(self, tmp_path):
        source = letter_docx(tmp_path / 'in.docx')
        ran = subprocess.run(
            [sys.executable, str(OFFICE / 'verify_bangla.py'), str(source)], capture_output=True, text=True
        )
        assert ran.returncode == 2
        assert json.loads(ran.stdout)['status'] == 'defects_found'
        fixed = tmp_path / 'out.docx'
        ran = subprocess.run(
            [sys.executable, str(OFFICE / 'fix_bangla.py'), str(source), '-o', str(fixed), '--new'],
            capture_output=True, text=True,
        )
        assert ran.returncode == 0, ran.stdout + ran.stderr
        assert json.loads(ran.stdout)['output'] == str(fixed)


def test_a_deliberately_larger_bangla_size_is_kept(tmp_path):
    """Offices set Bangla a size above the English beside it; the fix must not undo that."""
    source = tmp_path / 'in.docx'
    document = docx.Document()
    run = document.add_paragraph().add_run(SENTENCE)
    run.font.size = Pt(12)
    rpr = run._r.get_or_add_rPr()
    size_cs = etree.SubElement(rpr, f'{{{W}}}szCs')
    size_cs.set(f'{{{W}}}val', '28')
    document.save(source)
    out = tmp_path / 'out.docx'
    fix(source, out, new=False, embed=False)
    assert bangla_runs(out)[0].find('w:rPr/w:szCs', NS).get(f'{{{W}}}val') == '28'
    assert not [f for f in verify(out, source)['findings'] if 'szCs' in f['detail']]


def test_garbled_text_the_original_had_is_a_note_after_a_fix(tmp_path):
    """A DOCX made from a broken PDF layer carries (cid:N) and drawn-order Bangla. Fixing its fonts
    changes every run's formatting; the garbled text is still the original's, not the edit's."""
    source = tmp_path / 'in.docx'
    document = docx.Document()
    run = document.add_paragraph().add_run('বাংলােদশ িব(cid:10)িবদ(cid:11)ালয় ম(cid:14)ুরী কিমশন')
    run.font.size = Pt(10)
    document.save(source)
    out = tmp_path / 'out.docx'
    fix(source, out, new=False, embed=HAS_NIKOSH)
    result = verify(out, source)
    assert 'garbled' not in {f['check'] for f in result['findings']}
    assert 'garbled' in {f['check'] for f in result['notes']}


BIJOY = 'evsjv‡`k wek¦we`¨vjq gÄyix Kwgkb'
BIJOY_TEXT = bangla.compose('বাংলাদেশ বিশ্ববিদ্যালয় মঞ্জুরী কমিশন')
"""As written into a file: য় as one character (`bangla.compose`)."""


def bijoy_docx(path):
    """A letter typed in SutonnyMJ, as offices wrote them before Unicode: one paragraph whose word
    কি is split across two runs (`w` | `K`), a Bijoy paragraph style, and an English line."""
    document = docx.Document()
    style = document.styles.add_style('Bijoy Body', docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = 'SutonnyMJ'
    first = document.add_paragraph()
    for text in ('Avwg evsjvq Mvb MvB, ', 'w', 'K?'):
        run = first.add_run(text)
        run.font.name = 'SutonnyMJ'
        run.font.size = Pt(14)
    document.add_paragraph(BIJOY, style='Bijoy Body')
    document.add_paragraph('Memo No. 37.01.0000').runs[0].font.name = 'Calibri'
    document.save(path)
    return path


def docx_texts(path):
    return [p.text for p in docx.Document(path).paragraphs]


XLSX_MAIN = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
XLSX_PARTS = {
    '[Content_Types].xml': (
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        '<Override PartName="/xl/styles.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
        '<Override PartName="/xl/sharedStrings.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
        '</Types>'
    ),
    '_rels/.rels': (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Target="xl/workbook.xml" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"/>'
        '</Relationships>'
    ),
    'xl/workbook.xml': (
        f'<workbook xmlns="{XLSX_MAIN}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<sheets><sheet name="Sheet" sheetId="1" r:id="rId1"/></sheets></workbook>'
    ),
    'xl/_rels/workbook.xml.rels': (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Target="worksheets/sheet1.xml" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"/>'
        '<Relationship Id="rId2" Target="styles.xml" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles"/>'
        '<Relationship Id="rId3" Target="sharedStrings.xml" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings"/>'
        '</Relationships>'
    ),
    'xl/styles.xml': (
        f'<styleSheet xmlns="{XLSX_MAIN}">'
        '<fonts count="2"><font><sz val="11"/><name val="Arial"/></font>'
        '<font><sz val="12"/><name val="SutonnyMJ"/><charset val="0"/></font></fonts>'
        '<fills count="1"><fill><patternFill patternType="none"/></fill></fills>'
        '<borders count="1"><border/></borders>'
        '<cellStyleXfs count="1"><xf/></cellStyleXfs>'
        '<cellXfs count="2"><xf fontId="0"/><xf fontId="1" applyFont="1"/></cellXfs>'
        '</styleSheet>'
    ),
    'xl/sharedStrings.xml': (
        f'<sst xmlns="{XLSX_MAIN}" count="3" uniqueCount="2"><si><t>{BIJOY}</t></si>'
        '<si><r><rPr><rFont val="SutonnyMJ"/></rPr><t>evsjv</t></r></si></sst>'
    ),
    'xl/worksheets/sheet1.xml': (
        f'<worksheet xmlns="{XLSX_MAIN}"><sheetData>'
        '<row r="1"><c r="A1" s="1" t="s"><v>0</v></c><c r="B1" s="0" t="s"><v>0</v></c></row>'
        '<row r="2"><c r="A2" s="0" t="s"><v>1</v></c></row>'
        '</sheetData></worksheet>'
    ),
}


def shared_strings_xlsx(path):
    """A workbook the way Excel writes it: text in sharedStrings.xml, the SutonnyMJ string shown
    by A1 (SutonnyMJ) and B1 (Arial), and a rich-text SutonnyMJ run in an Arial cell (A2)."""
    with zipfile.ZipFile(path, 'w') as archive:
        for name, content in XLSX_PARTS.items():
            archive.writestr(name, '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>' + content)
    return path


class TestBijoy:
    def test_docx_text_becomes_unicode_in_nikosh(self, tmp_path):
        out = tmp_path / 'out.docx'
        result = fix(bijoy_docx(tmp_path / 'in.docx'), out, embed=False)
        assert result['bijoy'] == {'runs': 4, 'fonts': ['SutonnyMJ']}
        assert docx_texts(out) == [bangla.compose('আমি বাংলায় গান গাই, কি?'), BIJOY_TEXT, 'Memo No. 37.01.0000']
        for run in bangla_runs(out):
            fonts = run.find('w:rPr/w:rFonts', NS)
            assert {fonts.get(f'{{{W}}}{slot}') for slot in ('ascii', 'hAnsi', 'cs')} == {'Nikosh'}
            assert run.find('w:rPr/w:lang', NS).get(f'{{{W}}}bidi') == 'bn-BD'
            assert run.find('w:rPr/w:szCs', NS) is not None or run.find('w:rPr/w:sz', NS) is None

    def test_the_schema_validator_accepts_a_converted_docx(self, tmp_path):
        source = bijoy_docx(tmp_path / 'in.docx')
        out = tmp_path / 'out.docx'
        fix(source, out, embed=HAS_NIKOSH)
        checked = subprocess.run(
            [sys.executable, str(OFFICE / 'validate.py'), str(out), '--original', str(source)],
            capture_output=True, text=True,
        )
        assert checked.returncode == 0, checked.stdout + checked.stderr

    def test_english_keeps_its_font(self, tmp_path):
        out = tmp_path / 'out.docx'
        fix(bijoy_docx(tmp_path / 'in.docx'), out, embed=False)
        memo = docx.Document(out).paragraphs[2].runs[0]
        assert memo.font.name == 'Calibri'

    def test_a_converted_docx_passes_the_gate(self, tmp_path):
        source = bijoy_docx(tmp_path / 'in.docx')
        out = tmp_path / 'out.docx'
        fix(source, out, embed=HAS_NIKOSH)
        result = verify(out, source)
        if HAS_NIKOSH:
            assert result['status'] == 'clean', result
        assert not [f for f in result['findings'] if f['check'] in ('lost', 'garbled')]

    def test_bijoy_left_in_the_output_fails_even_when_the_original_had_it(self, tmp_path):
        source = bijoy_docx(tmp_path / 'in.docx')
        result = verify(source, source)
        bijoy_findings = [f for f in result['findings'] if 'Bijoy text in SutonnyMJ' in f['detail']]
        assert len(bijoy_findings) == 4
        assert BIJOY_TEXT in ' '.join(f['detail'] for f in bijoy_findings)

    def test_a_lost_bijoy_paragraph_is_reported_in_unicode(self, tmp_path):
        source = bijoy_docx(tmp_path / 'in.docx')
        out = tmp_path / 'out.docx'
        fix(source, out, embed=False)
        document = docx.Document(out)
        document.paragraphs[1]._p.getparent().remove(document.paragraphs[1]._p)
        document.save(out)
        lost = [f['detail'] for f in verify(out, source)['findings'] if f['check'] == 'lost']
        assert lost == [f'missing from the output: {BIJOY_TEXT!r}']

    def test_pptx(self, tmp_path):
        source = TestFixPptx().deck(tmp_path / 'in.pptx', text=BIJOY, font='SutonnyMJ')
        out = tmp_path / 'out.pptx'
        result = fix(source, out)
        assert result['bijoy'] == {'runs': 1, 'fonts': ['SutonnyMJ']}
        rpr = xml(out, 'ppt/slides/slide1.xml').find('.//a:r/a:rPr', NS)
        assert rpr.find('a:latin', NS).get('typeface') == 'Nikosh'
        assert rpr.find('a:cs', NS).get('typeface') == 'Nikosh'
        assert rpr.get('lang') == 'bn-BD'
        assert pptx.Presentation(out).slides[0].shapes[0].text_frame.text == BIJOY_TEXT
        assert verify(out, source)['status'] == 'clean'

    def workbook(self, path):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        bijoy_font, arial = openpyxl.styles.Font(name='SutonnyMJ', size=12), openpyxl.styles.Font(name='Arial')
        for cell, value, font in (
            ('A1', BIJOY, bijoy_font),
            ('A2', 2026, bijoy_font),
            ('B1', BIJOY, arial),
            ('C1', 'Total', arial),
        ):
            sheet[cell] = value
            sheet[cell].font = font
        sheet['A3'] = '=A2*2'
        workbook.save(path)
        return path

    def test_xlsx_cells_convert_and_numbers_stay_numbers(self, tmp_path):
        source = self.workbook(tmp_path / 'in.xlsx')
        out = tmp_path / 'out.xlsx'
        result = fix(source, out)
        assert result['bijoy'] == {'cells': 1, 'strings': 0, 'numbers': 1, 'fonts': ['SutonnyMJ']}
        sheet = openpyxl.load_workbook(out).active
        assert sheet['A1'].value == BIJOY_TEXT
        assert sheet['A1'].font.name == 'Nikosh'
        assert sheet['A2'].value == 2026
        assert sheet['A3'].value == '=A2*2'
        assert sheet['B1'].value == BIJOY
        assert sheet['C1'].value == 'Total'
        gate = verify(out, source)
        assert gate['status'] == 'clean', gate
        assert [note['where'] for note in gate['notes']] == ['Sheet!B1']

    def test_xlsx_parts_without_bijoy_are_untouched(self, tmp_path):
        source = self.workbook(tmp_path / 'in.xlsx')
        out = tmp_path / 'out.xlsx'
        fix(source, out)
        with zipfile.ZipFile(source) as before, zipfile.ZipFile(out) as after:
            changed = {name for name in before.namelist() if before.read(name) != after.read(name)}
        assert changed == {'xl/styles.xml', 'xl/worksheets/sheet1.xml'}

    def test_xlsx_shared_strings(self, tmp_path):
        """Excel stores text once in sharedStrings.xml. A string shown both in SutonnyMJ and in
        Arial keeps its text for the Arial cell and gets a converted copy for the Bijoy one."""
        source = shared_strings_xlsx(tmp_path / 'in.xlsx')
        out = tmp_path / 'out.xlsx'
        result = fix(source, out)
        assert result['bijoy'] == {'cells': 2, 'strings': 2, 'numbers': 0, 'fonts': ['SutonnyMJ']}
        sheet = openpyxl.load_workbook(out).active
        assert (sheet['A1'].value, sheet['B1'].value, sheet['A2'].value) == (BIJOY_TEXT, BIJOY, 'বাংলা')
        assert sheet['A1'].font.name == 'Nikosh'
        table = xml(out, 'xl/sharedStrings.xml')
        assert table.get('uniqueCount') == '3'
        assert verify(out, source)['status'] == 'clean'

    def test_xlsx_rich_text_runs(self, tmp_path):
        rich = pytest.importorskip('openpyxl.cell.rich_text')
        from openpyxl.cell.text import InlineFont

        source = tmp_path / 'in.xlsx'
        workbook = openpyxl.Workbook()
        workbook.active['A1'] = rich.CellRichText(
            rich.TextBlock(InlineFont(rFont='SutonnyMJ'), 'evsjv'),
            rich.TextBlock(InlineFont(rFont='Arial'), ' (Bangla)'),
        )
        workbook.save(source)
        out = tmp_path / 'out.xlsx'
        fix(source, out)
        cell = openpyxl.load_workbook(out, rich_text=True).active['A1'].value
        assert [(block.font.rFont, block.text) for block in cell] == [('Nikosh', 'বাংলা'), ('Arial', ' (Bangla)')]


SPLIT = 'বিশ্ববিদ্যাল\u09af\u09bc ও বড\u09bc'
"""য় and ড় as letter + nukta, the form NFC gives and Word draws with the dot beside the letter."""


class TestNukta:
    def test_compose_writes_single_characters(self):
        assert bangla.compose(SPLIT) == 'বিশ্ববিদ্যাল\u09df ও ব\u09dc'
        assert bangla.normalize(bangla.compose(SPLIT)) == bangla.normalize(SPLIT)

    def test_bijoy_output_is_composed_once_written(self, tmp_path):
        source = tmp_path / 'in.docx'
        document = docx.Document()
        document.add_paragraph().add_run('evsjvq').font.name = 'SutonnyMJ'
        document.save(source)
        out = tmp_path / 'out.docx'
        fix(source, out, embed=False)
        assert docx.Document(out).paragraphs[0].text == 'বাংলা\u09df'

    def test_fix_and_gate_on_a_split_docx(self, tmp_path):
        source = tmp_path / 'in.docx'
        document = docx.Document()
        document.add_paragraph().add_run(SPLIT).font.name = 'Nikosh'
        document.save(source)
        assert any('letter + nukta' in f['detail'] for f in verify(source)['findings'])
        out = tmp_path / 'out.docx'
        result = fix(source, out, embed=False)
        assert result['runs_changed']['nukta'] == 1
        assert '\u09bc' not in docx.Document(out).paragraphs[0].text
        assert not [f for f in verify(out, source)['findings'] if 'nukta' in f['detail']]

    def test_pptx_and_xlsx(self, tmp_path):
        deck = TestFixPptx().deck(tmp_path / 'in.pptx', text=SPLIT, font='Nikosh')
        fix(deck, tmp_path / 'out.pptx')
        assert '\u09bc' not in pptx.Presentation(tmp_path / 'out.pptx').slides[0].shapes[0].text_frame.text
        workbook = openpyxl.Workbook()
        workbook.active['A1'] = SPLIT
        workbook.save(tmp_path / 'in.xlsx')
        assert fix(tmp_path / 'in.xlsx', tmp_path / 'out.xlsx')['nukta'] == 1
        assert '\u09bc' not in openpyxl.load_workbook(tmp_path / 'out.xlsx').active['A1'].value
