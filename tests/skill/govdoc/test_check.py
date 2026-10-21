"""Tests for skill/govdoc/scripts/check.py on text drafts and built DOCX files.

The text tests are pure Python. The DOCX tests need python-docx and lxml:
    python -m pytest tests/skill/govdoc
"""
import importlib.util
import json
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[3] / 'skill' / 'govdoc' / 'scripts'))
import check  # noqa: E402

HAS_DOCX = importlib.util.find_spec('docx') is not None

MEMO = """গণপ্রজাতন্ত্রী বাংলাদেশ সরকার
শিক্ষা মন্ত্রণালয়
স্মারক নম্বর: ৩৭.০০.০০০০.০০০.০৭১.৯৯.০০১২.২৬.২১৪
তারিখ:
০৫ কার্তিক ১৪৩৩
২১ অক্টোবর ২০২৬
বিষয়: তথ্য হালনাগাদ প্রসঙ্গে।
উপর্যুক্ত বিষয়ের পরিপ্রেক্ষিতে জানানো যাচ্ছে যে, তথ্য হালনাগাদ করা প্রয়োজন।
২। এমতাবস্থায়, তথ্য প্রেরণের জন্য নির্দেশক্রমে অনুরোধ করা হলো।
(মোহাম্মদ রফিকুল ইসলাম)
উপসচিব
ফোন: ০২-৯৫৪৯৬২১
ই-মেইল: ict@moedu.gov.bd
"""


def run(tmp_path, text, kind=None, original=None, name='draft.txt'):
    path = tmp_path / name
    path.write_text(text, encoding='utf-8')
    original_path = None
    if original is not None:
        original_path = tmp_path / ('original-' + name)
        original_path.write_text(original, encoding='utf-8')
    return check.check(path, kind, original_path)


def checks(result, level='defects'):
    return [(f['check'], f['text']) for f in result[level]]


def test_clean_memo(tmp_path):
    result = run(tmp_path, MEMO, 'memo')
    assert result == {'status': 'clean', 'defects': [], 'notes': []}


def test_wrong_bangla_date(tmp_path):
    result = run(tmp_path, MEMO.replace('০৫ কার্তিক', '০৬ কার্তিক'), 'memo')
    assert checks(result) == [('date', '০৬ কার্তিক ১৪৩৩')]
    assert '০৫ কার্তিক ১৪৩৩' in result['defects'][0]['fix']


def test_spelling(tmp_path):
    result = run(tmp_path, MEMO.replace('প্রয়োজন।', 'প্রয়োজন। উপরোক্ত সরকারী নির্দেশাবলী মানতে হবে।'), 'memo')
    assert checks(result) == [('spelling', 'উপরোক্ত'), ('spelling', 'সরকারী'), ('spelling', 'নির্দেশাবলী')]


def test_spelling_is_matched_as_whole_words(tmp_path):
    result = run(tmp_path, MEMO.replace('প্রয়োজন।', 'প্রয়োজন। দেশীয় পণ্য।'), 'memo')
    assert result['status'] == 'clean'


def test_memo_has_no_salutation_or_courtesy_close(tmp_path):
    text = MEMO.replace('উপর্যুক্ত', 'মহোদয়,\nউপর্যুক্ত').replace('(মোহাম্মদ', 'আপনার বিশ্বস্ত,\n(মোহাম্মদ')
    assert checks(run(tmp_path, text, 'memo')) == [('structure', 'মহোদয়,'), ('structure', 'আপনার বিশ্বস্ত,')]


def test_letter_needs_its_courtesy_close(tmp_path):
    assert checks(run(tmp_path, MEMO, 'letter')) == [('structure', 'the closing "আপনার বিশ্বস্ত,"')]


def test_order_has_no_subject_and_needs_its_title(tmp_path):
    result = run(tmp_path, MEMO, 'order')
    assert ('structure', 'the title "অফিস আদেশ"') in checks(result)
    assert ('structure', 'বিষয়') in checks(result)


def test_missing_subject(tmp_path):
    result = run(tmp_path, MEMO.replace('বিষয়: তথ্য হালনাগাদ প্রসঙ্গে।\n', ''), 'notice')
    assert ('structure', 'the বিষয় line') in checks(result)


def test_notes_digits_register_contact_visarga(tmp_path):
    text = (MEMO.replace('তথ্য প্রেরণের', '15 দিনের মধ্যে তথ্য প্রেরণের')
                .replace('করা প্রয়োজন।', 'করা হইবে।')
                .replace('ফোন: ০২-৯৫৪৯৬২১\n', '')
                .replace('বিষয়:', 'বিষয়ঃ'))
    result = run(tmp_path, text, 'memo')
    assert result['status'] == 'clean'
    assert {c for c, _ in checks(result, 'notes')} == {'digits', 'register', 'contact', 'visarga'}
    assert ('digits', '15') in checks(result, 'notes')


def test_email_digits_are_not_flagged(tmp_path):
    result = run(tmp_path, MEMO.replace('ict@moedu.gov.bd', 'ict2026@moedu.gov.bd'), 'memo')
    assert result['notes'] == []


def test_meeting_before_its_notice(tmp_path):
    text = MEMO.replace('বিষয়:', 'সভার নোটিশ\nবিষয়:').replace(
        'তথ্য হালনাগাদ করা প্রয়োজন।', 'কমিটির সভা আগামী ১৬.০৫.২০২৬ তারিখ সকাল ১০:০০ টায় অনুষ্ঠিত হবে।')
    assert checks(run(tmp_path, text, 'meeting')) == [('date', '১৬.০৫.২০২৬')]
    later = text.replace('১৬.০৫.২০২৬', '০৫.১১.২০২৬')
    assert run(tmp_path, later, 'meeting')['status'] == 'clean'


def test_file_number_is_not_read_as_a_meeting_date(tmp_path):
    text = MEMO.replace('বিষয়:', 'সভার নোটিশ\nবিষয়:').replace(
        'তথ্য হালনাগাদ করা প্রয়োজন।', 'সূত্র ০৪.০০.০০০০.০০০.৪১৬.৯৯.০০০৭.২৬.৪৬৫ অনুযায়ী আগামী ০৫.১১.২০২৬ তারিখ সভা।')
    assert run(tmp_path, text, 'meeting')['status'] == 'clean'


def test_faults_of_the_original_are_notes(tmp_path):
    original = MEMO.replace('প্রয়োজন।', 'প্রয়োজন। সরকারী।')
    edited = original.replace('তথ্য প্রেরণের', 'হালনাগাদ তথ্য প্রেরণের')
    result = run(tmp_path, edited, 'memo', original=original)
    assert result['status'] == 'clean'
    assert checks(result, 'notes') == [('spelling', 'সরকারী')]


def test_rejects_unknown_type_and_file(tmp_path):
    with pytest.raises(ValueError):
        run(tmp_path, MEMO, 'gazette')
    with pytest.raises(ValueError):
        run(tmp_path, MEMO, None, name='draft.pdf')


def test_cli_exit_codes(tmp_path, capsys, monkeypatch):
    path = tmp_path / 'draft.txt'
    path.write_text(MEMO.replace('সরকার\n', 'সরকারী\n', 1), encoding='utf-8')
    monkeypatch.setattr(sys, 'argv', ['check.py', str(path), '--type', 'memo'])
    assert check.main() == 2
    assert json.loads(capsys.readouterr().out)['defects'][0]['check'] == 'spelling'


@pytest.mark.skipif(not HAS_DOCX, reason='python-docx is not installed')
def test_reads_docx_tables_in_order(tmp_path):
    from docx import Document
    document = Document()
    document.add_paragraph('বিষয়: পরীক্ষা।')
    table = document.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ('নম্বর: ১২', 'তারিখ:', '০৬ কার্তিক ১৪৩৩')):
        cell.text = text
    table.rows[0].cells[2].add_paragraph('২১ অক্টোবর ২০২৬')
    path = tmp_path / 'memo.docx'
    document.save(path)
    assert check.paragraphs(path)[:5] == ['বিষয়: পরীক্ষা।', 'নম্বর: ১২', 'তারিখ:', '০৬ কার্তিক ১৪৩৩', '২১ অক্টোবর ২০২৬']
    assert checks(check.check(path)) == [('date', '০৬ কার্তিক ১৪৩৩')]
