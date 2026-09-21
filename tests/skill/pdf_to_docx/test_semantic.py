"""Tests for the semantic-editable route: skill/pdf-to-docx/scripts/semantic.py.

The documents this mode exists for are official orders and memos — a crest, headings, mixed Bangla
and English, a ruled table, a numbered list, a signature and a QR code. What it promises is that
the text edits like a Word document (paragraphs that reflow, not positioned boxes) while every
image stays exactly as the PDF drew it.

Run with pdfplumber, python-docx, Pillow, reportlab and qrcode installed:
    python -m pytest tests/skill/pdf_to_docx/test_semantic.py
"""
import shutil
import sys
import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn

from helpers import (
    HEIGHT,
    SCRIPTS,
    bengali_font,
    build_pdf,
    load,
    make_logo,
    make_qr,
    make_signature,
)

sys.path.insert(0, str(SCRIPTS))
probe = load('probe')
semantic = load('semantic')
convert = load('convert')
verify = load('verify_conversion')

needs_soffice = pytest.mark.skipif(
    shutil.which('soffice') is None, reason='LibreOffice is not installed'
)

QR_PAYLOAD = 'BDREN/OFFICE/ORDER/2026/0917'
BANGLA_LINE = 'বাংলাদেশ গবেষণা ও শিক্ষা নেটওয়ার্ক'
BODY = [
    'The competent authority has been pleased to approve the deployment of the document',
    'conversion service for all member universities with immediate effect.',
]
STEPS = [
    '1. Nominate a focal point for each member university.',
    '2. Verify that every converted order opens in Word.',
    '3. Report completion to the Admin section by 30 September 2026.',
]
TABLE = [
    ['Phase', 'Scope', 'Date'],
    ['Phase 0', 'Pilot with two universities', '20 Sep 2026'],
    ['Phase 1', 'All member universities', '30 Sep 2026'],
]
COLUMNS = (72, 200, 400, 523)
ROW_HEIGHT = 22


@pytest.fixture(scope='module')
def assets(tmp_path_factory):
    directory = tmp_path_factory.mktemp('assets')
    return {
        'logo': make_logo(directory / 'logo.png'),
        'signature': make_signature(directory / 'sign.png'),
        'qr': make_qr(directory / 'qr.png', QR_PAYLOAD),
    }


def draw_order(page, assets):
    """One page that carries every element an official order does."""
    page.drawImage(str(assets['logo']), 72, HEIGHT - 130, width=54, height=54)
    page.setFont('Helvetica-Bold', 17)
    page.drawString(140, HEIGHT - 100, 'Bangladesh Research and Education Network')
    bangla = bengali_font()
    if bangla:
        page.setFont(bangla, 12)
        page.drawString(140, HEIGHT - 120, BANGLA_LINE)
    page.setFont('Helvetica', 10)
    page.drawString(140, HEIGHT - 140, 'Memo No: BdREN/ADMIN/2026/117      Date: 17 September 2026')

    page.setFont('Helvetica-Bold', 14)
    page.drawString(72, HEIGHT - 185, 'Office Order')
    page.setFont('Helvetica', 11)
    text = page.beginText(72, HEIGHT - 215)
    for line in BODY:
        text.textLine(line)
    page.drawText(text)

    top = HEIGHT - 260
    for index, row in enumerate(TABLE):
        y = top - index * ROW_HEIGHT
        page.line(COLUMNS[0], y, COLUMNS[-1], y)
        page.line(COLUMNS[0], y - ROW_HEIGHT, COLUMNS[-1], y - ROW_HEIGHT)
        page.setFont('Helvetica-Bold' if index == 0 else 'Helvetica', 10)
        for column, cell in enumerate(row):
            page.drawString(COLUMNS[column] + 5, y - 15, cell)
    bottom = top - len(TABLE) * ROW_HEIGHT
    for x in COLUMNS:
        page.line(x, top, x, bottom)

    page.setFont('Helvetica', 11)
    steps = page.beginText(72, bottom - 30)
    for line in STEPS:
        steps.textLine(line)
    page.drawText(steps)

    page.drawImage(str(assets['signature']), 72, bottom - 140, width=130, height=45)
    page.setFont('Helvetica', 10)
    page.drawString(72, bottom - 155, 'Director (Administration)')
    page.drawImage(str(assets['qr']), 430, bottom - 155, width=80, height=80)


@pytest.fixture(scope='module')
def order(assets, tmp_path_factory):
    path = tmp_path_factory.mktemp('pdf') / 'order.pdf'
    return build_pdf(path, [lambda page: draw_order(page, assets)])


def draw_continuation(page, assets):
    """A second page ending in the same signature block as the first."""
    page.setFont('Helvetica', 11)
    page.drawString(72, HEIGHT - 100, 'Annexure: the rollout completes on 30 September 2026.')
    page.drawImage(str(assets['signature']), 72, HEIGHT - 220, width=130, height=45)
    page.setFont('Helvetica', 10)
    page.drawString(72, HEIGHT - 235, 'Director (Administration)')


@pytest.fixture(scope='module')
def repeated(assets, tmp_path_factory):
    """Two pages that share one signature raster, converted."""
    directory = tmp_path_factory.mktemp('repeated')
    source = build_pdf(
        directory / 'two.pdf',
        [lambda page: draw_order(page, assets), lambda page: draw_continuation(page, assets)],
    )
    output = directory / 'two.docx'
    semantic.build(source, output)
    return source, output


@pytest.fixture(scope='module')
def built(order, tmp_path_factory):
    output = tmp_path_factory.mktemp('docx') / 'order.docx'
    return output, semantic.build(order, output)


LETTERHEAD = 'Office of the Registrar, Establishment-01'
FOOTER_MEMO = 'Memo No: KU/ADMIN/2026/117'
STATEMENT = [
    ['Sl', 'Discipline', 'Name', 'Amount (BDT)'],
    ['1', 'Drawing and Painting', 'Officer 1', '12,137'],
    ['2', 'Drawing and Painting', 'Officer 2', '12,274'],
]
STATEMENT_COLUMNS = (50, 80, 220, 340, 470)


def draw_running_page(page, assets, number, total, rows):
    """A page of a report: letterhead above, page number below, a ruled table between."""
    page.drawImage(str(assets['logo']), 50, HEIGHT - 78, width=36, height=36)
    page.setFont('Helvetica-Bold', 13)
    page.drawString(96, HEIGHT - 56, 'Khulna University')
    page.setFont('Helvetica', 9)
    page.drawString(96, HEIGHT - 68, LETTERHEAD)
    page.line(50, HEIGHT - 88, 470, HEIGHT - 88)

    y = HEIGHT - 130
    for index, row in enumerate(rows):
        page.setFont('Helvetica-Bold' if index == 0 else 'Helvetica', 9)
        page.line(STATEMENT_COLUMNS[0], y, STATEMENT_COLUMNS[-1], y)
        for column, cell in enumerate(row):
            page.drawString(STATEMENT_COLUMNS[column] + 4, y - 12, cell)
        y -= 16
    page.line(STATEMENT_COLUMNS[0], y, STATEMENT_COLUMNS[-1], y)
    for x in STATEMENT_COLUMNS:
        page.line(x, HEIGHT - 130, x, y)

    page.line(50, 60, 470, 60)
    page.setFont('Helvetica', 8)
    page.drawString(50, 48, FOOTER_MEMO)
    page.drawRightString(470, 48, f'Page {number} of {total}')


@pytest.fixture(scope='module')
def report(assets, tmp_path_factory):
    """Two pages sharing a letterhead and a numbered footer, converted."""
    directory = tmp_path_factory.mktemp('report')
    source = build_pdf(
        directory / 'report.pdf',
        [
            lambda page: draw_running_page(page, assets, 1, 2, STATEMENT),
            lambda page: draw_running_page(page, assets, 2, 2, STATEMENT[:1] + STATEMENT[1:]),
        ],
    )
    output = directory / 'report.docx'
    return source, output, semantic.build(source, output)


def part_text(part):
    return '\n'.join(paragraph.text for paragraph in part.paragraphs)


def document_xml(path):
    return zipfile.ZipFile(path).read('word/document.xml').decode('utf8')


def paragraph_texts(path):
    return [paragraph.text for paragraph in Document(str(path)).paragraphs]


class TestSemantic:
    def test_it_recovers_the_text_the_pdf_carries(self, built, order):
        output, summary = built
        recovered = probe.normalize(probe.docx_text(output))
        source = probe.normalize(probe.pdf_text(order))
        assert len(recovered) >= 0.9 * len(source)
        assert 'competent authority' in probe.docx_text(output)
        assert summary['paragraphs'] > 0

    def test_the_text_is_in_paragraphs_that_reflow(self, built):
        """The defect this mode exists to remove: a line per positioned text box."""
        output, _ = built
        assert '<w:txbxContent>' not in document_xml(output)
        assert any('competent authority' in text for text in paragraph_texts(output))

    def test_the_headings_come_back_as_headings(self, built):
        output, summary = built
        styles = [p.style.name for p in Document(str(output)).paragraphs]
        assert summary['headings'] >= 1
        assert any(name.startswith('Heading') for name in styles)

    def test_the_numbered_steps_come_back_as_a_list(self, built):
        output, summary = built
        document = Document(str(output))
        items = [p.text for p in document.paragraphs if p.style.name == 'List Number']
        assert summary['list_items'] == len(STEPS)
        assert len(items) == len(STEPS)
        assert items[0].startswith('Nominate a focal point')

    def test_the_ruled_grid_comes_back_as_a_table(self, built):
        output, summary = built
        tables = Document(str(output)).tables
        assert summary['tables'] == 1
        assert len(tables) == 1
        assert len(tables[0].rows) == len(TABLE)
        assert [cell.text for cell in tables[0].rows[0].cells] == TABLE[0]
        assert tables[0].rows[2].cells[2].text == '30 Sep 2026'

    def test_table_text_is_not_repeated_as_paragraphs(self, built):
        """Words inside a table region must be spent once, or every cell is written twice."""
        output, _ = built
        outside = '\n'.join(paragraph_texts(output))
        assert 'Pilot with two universities' not in outside

    def test_every_picture_survives(self, built, order):
        output, summary = built
        assert summary['pictures'] == probe.pdf_image_count(order) == 3
        assert len(probe.docx_media(output)) == 3

    def test_no_picture_stands_in_for_the_page(self, built):
        """A crop that swallowed the page would pass the image count and lose the text."""
        output, _ = built
        width, height = probe.docx_page_size(output)
        assert all(cx < 0.9 * width and cy < 0.9 * height for cx, cy in probe.docx_extents(output))

    def test_the_qr_still_decodes(self, built, order):
        output, _ = built
        findings = []
        verify._check_qr(findings, probe.pdf_media(order), probe.docx_media(output))
        assert findings == []

    def test_the_page_keeps_its_size(self, built, order):
        output, _ = built
        source = probe.pdf_page_size(order)
        assert probe.docx_page_size(output) == pytest.approx(source, abs=1.0)

    def test_a_block_of_short_lines_stays_a_stack_of_lines(self):
        """A letterhead is lines that never reached the edge; prose is lines that wrapped at it."""
        frame = {'left': 72.0, 'right': 523.0}
        lines = [{'x0': 96.0, 'x1': 258.0}, {'x0': 110.0, 'x1': 241.0}]
        assert semantic._is_stacked({'x0': 96.0, 'x1': 258.0, 'lines': lines}, frame) is True
        wrapped = [{'x0': 72.0, 'x1': 520.0}, {'x0': 72.0, 'x1': 388.0}]
        assert semantic._is_stacked({'x0': 72.0, 'x1': 520.0, 'lines': wrapped}, frame) is False

    def test_a_signer_block_against_the_right_margin_is_still_a_stack(self):
        """Name and designation reach the right edge without filling a line; joined, they read as
        one sentence."""
        frame = {'left': 72.0, 'right': 523.0}
        signer = [{'x0': 440.0, 'x1': 520.0}, {'x0': 425.0, 'x1': 521.0}]
        assert semantic._is_stacked({'x0': 425.0, 'x1': 521.0, 'lines': signer}, frame) is True

    def test_it_refuses_a_page_with_nothing_on_it(self, tmp_path):
        blank = build_pdf(tmp_path / 'blank.pdf', [lambda page: None])
        with pytest.raises(semantic.SemanticError, match='layout-editable'):
            semantic.build(blank, tmp_path / 'blank.docx')


@pytest.mark.skipif(bengali_font() is None, reason='no Bengali font installed')
class TestBangla:
    def test_the_bangla_line_survives(self, built):
        output, _ = built
        assert any('ঀ' <= char <= '৿' for char in probe.docx_text(output))

    def test_a_bangla_run_names_its_font_in_the_complex_script_slot(self, built):
        """Word picks the Bengali face from w:cs/w:eastAsia; the Latin slot alone leaves boxes."""
        output, _ = built
        for paragraph in Document(str(output)).paragraphs:
            if not any('ঀ' <= char <= '৿' for char in paragraph.text):
                continue
            fonts = paragraph.runs[0]._element.rPr.rFonts
            assert fonts.get(qn('w:cs'))
            assert fonts.get(qn('w:eastAsia'))
            return
        pytest.fail('no Bangla paragraph in the output')


class TestRunningBands:
    def test_the_letterhead_becomes_a_real_header(self, report):
        _, output, summary = report
        header = Document(str(output)).sections[0].header
        assert summary['running_header'] is True
        assert LETTERHEAD in part_text(header)
        assert 'graphic' in header.paragraphs[0]._p.xml or any(
            'graphic' in paragraph._p.xml for paragraph in header.paragraphs
        )

    def test_the_letterhead_is_not_left_in_the_body_as_well(self, report):
        """Repeated in the body it stops repeating the moment the text reflows."""
        _, output, _ = report
        assert LETTERHEAD not in '\n'.join(paragraph_texts(output))

    def test_the_footer_counts_its_own_pages(self, report):
        _, output, summary = report
        footer = Document(str(output)).sections[0].footer
        assert summary['running_footer'] is True
        assert FOOTER_MEMO in part_text(footer)
        assert 'PAGE' in footer.paragraphs[0]._p.xml

    def test_a_reference_number_is_not_turned_into_a_page_number(self, report):
        """`Establishment-01` is the office's own number; a PAGE field there rewrites it."""
        _, output, _ = report
        header = Document(str(output)).sections[0].header
        assert 'Establishment-01' in part_text(header)
        assert 'PAGE' not in header.paragraphs[0]._p.xml

    def test_the_gate_counts_a_running_band_once_per_page(self, report):
        source, output, _ = report
        result = verify.check(source, output, mode='semantic-editable')
        assert result['findings'] == []
        assert result['text_ratio'] >= 0.95

    def test_a_field_set_against_the_right_margin_is_tabbed_there(self, report):
        """Left-aligning it at its own start pushes the last word past the margin and wraps it."""
        from docx.enum.text import WD_TAB_ALIGNMENT

        _, output, _ = report
        footer = Document(str(output)).sections[0].footer
        stops = [
            stop
            for paragraph in footer.paragraphs
            for stop in paragraph.paragraph_format.tab_stops
        ]
        assert any(stop.alignment == WD_TAB_ALIGNMENT.RIGHT for stop in stops)

    def test_the_table_keeps_the_column_widths_the_page_gave_it(self, report):
        _, output, _ = report
        table = Document(str(output)).tables[0]
        expected = [
            STATEMENT_COLUMNS[index + 1] - STATEMENT_COLUMNS[index]
            for index in range(len(STATEMENT_COLUMNS) - 1)
        ]
        widths = [column.width.pt for column in table.columns]
        assert widths == pytest.approx(expected, abs=1.0)


class TestGate:
    def test_the_gate_passes_a_semantic_conversion(self, built, order):
        output, _ = built
        result = verify.check(order, output, mode='semantic-editable')
        assert result['findings'] == []
        assert result['status'] == 'clean'

    def test_the_gate_fails_a_layout_conversion_in_semantic_mode(self, built, order, tmp_path):
        """layout-editable puts every line in a box, which is exactly what this mode forbids."""
        output, _ = built
        boxed = tmp_path / 'boxed.docx'
        _box_the_text(output, boxed)
        result = verify.check(order, boxed, mode='semantic-editable')
        assert 'reflows' in [finding['check'] for finding in result['findings']]

    @needs_soffice
    def test_the_gate_counts_the_pages_a_reader_sees(self, built, order, tmp_path):
        """A reflowed document spills onto a page its own breaks know nothing about, which is how
        a two-page order was delivered as three and still counted as two."""
        output, _ = built
        spilled = tmp_path / 'spilled.docx'
        _add_page_breaks(output, spilled, 2)
        assert probe.docx_pages(spilled) == 3
        assert probe.rendered_pages(spilled) >= 3
        result = verify.check(order, spilled, mode='semantic-editable')
        assert 'pages' in [finding['check'] for finding in result['findings']]

    def test_the_gate_refuses_a_page_set_in_columns(self, tmp_path):
        """Lines are recovered by their vertical position, so two columns come back interleaved."""
        def draw(page):
            page.setFont('Helvetica', 10)
            for index in range(12):
                page.drawString(60, HEIGHT - 100 - index * 14, f'LEFT line {index} of the column.')
                page.drawString(320, HEIGHT - 100 - index * 14, f'RIGHT line {index} of it.')

        source = build_pdf(tmp_path / 'twocol.pdf', [draw])
        output = tmp_path / 'twocol.docx'
        semantic.build(source, output)
        result = verify.check(source, output, mode='semantic-editable')
        assert 'columns' in [finding['check'] for finding in result['findings']]

    def test_the_gate_does_not_call_a_table_a_column_layout(self, report):
        """A table's cells leave gaps as wide as a gutter; reading across one is correct."""
        source, output, _ = report
        result = verify.check(source, output, mode='semantic-editable')
        assert 'columns' not in [finding['check'] for finding in result['findings']]

    def test_the_gate_still_catches_a_lost_logo(self, built, order, tmp_path):
        """Re-rendered pictures no longer match by pixel size, so shape has to carry the check."""
        output, _ = built
        stripped = tmp_path / 'stripped.docx'
        _drop_one_picture(output, stripped)
        result = verify.check(order, stripped, mode='semantic-editable')
        assert 'images' in [finding['check'] for finding in result['findings']]

    def test_pictures_are_matched_by_the_rectangle_they_fill(self, built, order):
        """Pixel counts cannot carry the check: a picture may be re-rendered, and Word stores one
        copy of a raster placed twice. The rectangle is what both sides agree on."""
        output, _ = built
        drawn = sorted((p['width'], p['height']) for p in probe.docx_pictures(output))
        placed = sorted((entry['width'], entry['height']) for entry in probe.pdf_placements(order))
        assert len(drawn) == len(placed)
        for (width, height), (source_width, source_height) in zip(drawn, placed):
            assert width == pytest.approx(source_width, abs=1.0)
            assert height == pytest.approx(source_height, abs=1.0)

    def test_a_signature_on_every_page_is_not_a_lost_picture(self, repeated):
        """Word stores one part for a raster placed twice, so counting `word/media` under-counts
        the pictures — the defect that failed a clean conversion of a real two-page order."""
        source, output = repeated
        assert len(probe.docx_media(output)) < len(probe.pdf_placements(source))
        assert verify.check(source, output, mode='semantic-editable')['findings'] == []


def _box_the_text(source, target):
    """A document whose text sits in text boxes — what layout-editable produces."""
    with zipfile.ZipFile(source) as archive, zipfile.ZipFile(target, 'w') as output:
        for name in archive.namelist():
            payload = archive.read(name)
            if name == 'word/document.xml':
                payload = payload.replace(b'<w:body>', b'<w:body><w:txbxContent></w:txbxContent>')
            output.writestr(name, payload)
    return target


def _add_page_breaks(source, target, count):
    """The same document, laid out over more pages than the PDF had."""
    from docx.enum.text import WD_BREAK

    document = Document(str(source))
    for _ in range(count):
        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        document.add_paragraph('continued')
    document.save(str(target))
    return target


def _drop_one_picture(source, target):
    """A conversion that lost a logo: one media part removed, the anchor left behind."""
    dropped = False
    with zipfile.ZipFile(source) as archive, zipfile.ZipFile(target, 'w') as output:
        for name in archive.namelist():
            if name.startswith('word/media/') and not dropped:
                dropped = True
                continue
            output.writestr(name, archive.read(name))
    assert dropped
    return target
