"""Tests for skill/pdf-to-docx/scripts/.

Kept outside skill/pdf-to-docx/ on purpose: every file in a deployment skill directory is uploaded
into each user's code sandbox.

The conversion tests need LibreOffice on PATH and are skipped without it. Run with pypdf,
pdfplumber, python-docx, Pillow, opencv-python-headless, reportlab and qrcode installed:
    python -m pytest tests/skill/pdf_to_docx
"""
import json
import shutil
import subprocess
import sys
import zipfile

import pytest
import qrcode
from docx import Document
from docx.shared import Pt
from reportlab.lib.utils import ImageReader

from helpers import HEIGHT, SCRIPTS, WIDTH, build_pdf, load

QR_PAYLOAD = 'BDREN/OFFICE/ORDER/2026/0917'
BODY = [
    'Subject: Office order regarding the deployment of the document conversion service.',
    '',
    'The competent authority has been pleased to approve the deployment of the PDF to Word',
    'conversion route for all member universities with immediate effect. The service preserves',
    'the text of the source document so that recipients may edit, search and copy from it.',
    '',
    'All concerned are requested to take necessary action accordingly.',
]


sys.path.insert(0, str(SCRIPTS))
probe = load('probe')
restore = load('restore')
convert = load('convert')
verify = load('verify_conversion')

needs_soffice = pytest.mark.skipif(
    shutil.which('soffice') is None, reason='LibreOffice is not installed'
)


def draw_letter(page, with_qr=True):
    page.setFont('Helvetica-Bold', 18)
    page.drawString(72, HEIGHT - 90, 'Bangladesh Research and Education Network')
    page.setFont('Helvetica', 11)
    page.drawString(72, HEIGHT - 120, 'Memo No: BdREN/ADMIN/2026/117      Date: 17 September 2026')
    text = page.beginText(72, HEIGHT - 160)
    text.setFont('Helvetica', 11)
    for line in BODY:
        text.textLine(line)
    page.drawText(text)
    if with_qr:
        image = ImageReader(qrcode.make(QR_PAYLOAD).convert('RGB'))
        page.drawImage(image, 430, 690, width=90, height=90)


def draw_dark_page(page):
    page.setFillColorRGB(0.07, 0.1, 0.2)
    page.rect(0, 0, WIDTH, HEIGHT, stroke=0, fill=1)
    page.setFillColorRGB(1, 1, 1)
    page.setFont('Helvetica-Bold', 24)
    page.drawString(72, HEIGHT - 200, 'Annexure: Rollout Timeline')
    page.setFont('Helvetica', 12)
    page.drawString(72, HEIGHT - 240, 'Phase 0 completes on 30 September 2026.')


@pytest.fixture(scope='module')
def letter(tmp_path_factory):
    """One page: text and a QR code."""
    return build_pdf(tmp_path_factory.mktemp('pdf') / 'letter.pdf', [draw_letter])


@pytest.fixture(scope='module')
def order(tmp_path_factory):
    """Two pages: a letter with a QR, then a dark page with white text."""
    return build_pdf(tmp_path_factory.mktemp('pdf') / 'order.pdf', [draw_letter, draw_dark_page])


@pytest.fixture(scope='module')
def converted(order, tmp_path_factory):
    if shutil.which('soffice') is None:
        pytest.skip('LibreOffice is not installed')
    output = tmp_path_factory.mktemp('docx') / 'order.docx'
    return output, convert.convert(order, output)


def page_image_docx(path, pdf_path):
    """The failure this skill exists to prevent: every page flattened to a picture."""
    from docx.enum.text import WD_BREAK
    from pdf2image import convert_from_path

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Pt(WIDTH), Pt(HEIGHT)
    for number, page in enumerate(convert_from_path(str(pdf_path), dpi=72), start=1):
        image = path.parent / f'page{number}.png'
        page.save(image)
        run = document.add_paragraph().add_run()
        if number > 1:
            run.add_break(WD_BREAK.PAGE)
        run.add_picture(str(image), width=Pt(WIDTH), height=Pt(HEIGHT))
    document.save(str(path))
    return path


def checks(result):
    return [finding['check'] for finding in result['findings']]


def run(*args):
    return subprocess.run(
        [sys.executable, str(SCRIPTS / 'verify_conversion.py'), *map(str, args)],
        capture_output=True,
        text=True,
    )


class TestProbe:
    def test_reads_the_pdf_it_is_given(self, order):
        assert probe.pdf_pages(order) == 2
        width, height = probe.pdf_page_size(order)
        assert (round(width), round(height)) == (round(WIDTH), round(HEIGHT))
        assert 'competent authority' in probe.pdf_text(order)

    def test_finds_the_embedded_qr(self, letter):
        payloads = set()
        for entry in probe.pdf_media(letter):
            payloads |= probe.decode_qr(entry['image'])
        assert payloads == {QR_PAYLOAD}

    @needs_soffice
    def test_counts_each_character_once(self, converted, order):
        output, _ = converted
        source = probe.normalize(probe.pdf_text(order))
        delivered = probe.normalize(probe.docx_text(output))
        # LibreOffice writes every text box twice; counting the fallback would double this.
        assert len(delivered) == pytest.approx(len(source), rel=0.05)


class TestConvert:
    @needs_soffice
    def test_layout_editable_recovers_the_text_and_the_qr(self, converted, order):
        output, summary = converted
        assert summary['mode'] == 'layout-editable'
        assert summary['pages'] == summary['source_pages'] == 2
        assert summary['characters'] >= 0.7 * summary['source_characters']
        assert summary['images'] == summary['source_images'] == 1
        assert verify.check(order, output)['status'] == 'clean'

    @needs_soffice
    def test_layout_editable_embeds_no_page_image(self, converted):
        output, _ = converted
        width, height = probe.docx_page_size(output)
        assert all(
            cx < 0.9 * width or cy < 0.9 * height for cx, cy in probe.docx_extents(output)
        )

    @needs_soffice
    def test_the_dark_page_keeps_its_background_behind_the_text(self, converted):
        output, summary = converted
        assert summary['backgrounds_sent_to_back'] + summary['backgrounds_added'] == 1
        document = zipfile.ZipFile(output).read('word/document.xml').decode('utf8')
        assert 'mso-position-vertical-relative:page' in document
        assert 'z-index:-' in document

    def test_semantic_editable_reflows_a_text_only_pdf(self, tmp_path):
        source = build_pdf(tmp_path / 'plain.pdf', [lambda page: draw_letter(page, with_qr=False)])
        output = tmp_path / 'semantic.docx'
        summary = convert.convert(source, output, mode='semantic-editable')
        assert summary['pages'] == 1
        assert summary['characters'] >= 0.7 * summary['source_characters']
        assert 'competent authority' in ' '.join(p.text for p in Document(str(output)).paragraphs)

    def test_semantic_editable_keeps_the_pictures_of_an_illustrated_pdf(self, order, tmp_path):
        output = tmp_path / 'semantic.docx'
        summary = convert.convert(order, output, mode='semantic-editable')
        assert summary['pictures'] == summary['source_images'] == 1
        assert '<w:txbxContent>' not in zipfile.ZipFile(output).read('word/document.xml').decode()

    def test_visual_fidelity_makes_a_picture_of_every_page(self, order, tmp_path):
        output = tmp_path / 'visual.docx'
        summary = convert.convert(order, output, mode='visual-fidelity')
        assert summary['pages'] == summary['images'] == 2
        assert summary['characters'] == 0

    def test_it_refuses_a_mode_it_does_not_have(self, order, tmp_path):
        with pytest.raises(convert.ConversionError):
            convert.convert(order, tmp_path / 'out.docx', mode='pixel-perfect')

    def test_it_refuses_a_missing_input(self, tmp_path):
        with pytest.raises(convert.ConversionError):
            convert.convert(tmp_path / 'absent.pdf', tmp_path / 'out.docx')

    def test_it_refuses_an_output_that_is_not_a_docx(self, order, tmp_path):
        with pytest.raises(convert.ConversionError):
            convert.convert(order, tmp_path / 'out.pdf')


class TestRestore:
    def test_it_finds_the_full_bleed_fill(self, order, letter):
        assert restore.page_backgrounds(order) == {2: '121A33'}
        assert restore.page_backgrounds(letter) == {}

    def test_it_reads_colours_from_every_space(self):
        assert restore.srgb((0, 0, 0)) == '000000'
        assert restore.srgb((1,)) == 'FFFFFF'
        assert restore.srgb(0.5) == '808080'
        assert restore.srgb((0, 0, 0, 0)) == 'FFFFFF'

    @needs_soffice
    def test_it_reports_nothing_missing_when_nothing_is(self, converted, order):
        output, _ = converted
        assert restore.missing_images(order, output) == []

    @needs_soffice
    def test_it_re_inserts_an_image_the_import_dropped(self, letter, tmp_path):
        output = tmp_path / 'letter.docx'
        convert.convert(letter, output, repair=False)
        stripped = _without_media(output, tmp_path / 'stripped.docx')
        assert len(restore.missing_images(letter, stripped)) == 1

        summary = restore.restore(letter, stripped)
        assert summary['images_restored'] == 1
        assert restore.missing_images(letter, stripped) == []
        assert verify.check(letter, stripped)['status'] == 'clean'

    @needs_soffice
    def test_it_adds_a_background_the_import_really_dropped(self, order, tmp_path):
        output = tmp_path / 'order.docx'
        convert.convert(order, output, repair=False)
        stripped = _without_page_shapes(output, tmp_path / 'flat.docx')
        summary = restore.restore(order, stripped)
        assert summary['backgrounds_added'] == 1
        assert 'srgbClr val="121A33"' in zipfile.ZipFile(stripped).read('word/document.xml').decode()


def _without_media(source, target):
    """A conversion that lost its pictures: the media parts removed, the anchors left behind."""
    with zipfile.ZipFile(source) as archive, zipfile.ZipFile(target, 'w') as output:
        for name in archive.namelist():
            if not name.startswith('word/media/'):
                output.writestr(name, archive.read(name))
    return target


def _without_page_shapes(source, target):
    """A conversion that lost its page fill: every full-bleed VML shape removed."""
    from xml.etree import ElementTree

    with zipfile.ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    for prefix, uri in restore.NAMESPACES.items():
        ElementTree.register_namespace(prefix, uri)
    root = ElementTree.fromstring(parts['word/document.xml'])
    width, height = probe.docx_page_size(source)
    for parent in root.iter():
        for shape in list(parent):
            if (
                restore._is_vml_shape(shape)
                and restore._is_background_fill(shape)
                and restore._covers_page(shape, width, height)
            ):
                parent.remove(shape)
    parts['word/document.xml'] = ElementTree.tostring(root, encoding='UTF-8', xml_declaration=True)
    with zipfile.ZipFile(target, 'w') as output:
        for name, payload in parts.items():
            output.writestr(name, payload)
    return target


class TestGate:
    @needs_soffice
    def test_a_good_conversion_is_clean(self, converted, order):
        output, _ = converted
        result = verify.check(order, output)
        assert result['status'] == 'clean'
        assert result['findings'] == []
        assert result['text_ratio'] >= 0.7

    def test_a_page_image_document_fails_every_way_it_should(self, order, tmp_path):
        output = page_image_docx(tmp_path / 'images.docx', order)
        result = verify.check(order, output)
        assert result['status'] == 'defects'
        assert set(checks(result)) >= {'text', 'picture_book', 'images'}

    def test_a_hand_rebuilt_document_fails_however_good_its_text_is(self, order, tmp_path):
        """The 2026-09-18 failure: a styled python-docx rebuild, every image silently gone.

        Its text is the PDF's, so every text check passes; only the image check catches it. It
        must fail in the reflow mode too, or the mode becomes a licence to drop the seal.
        """
        output = tmp_path / 'rebuilt.docx'
        document = Document()
        section = document.sections[0]
        section.page_width, section.page_height = Pt(WIDTH), Pt(HEIGHT)
        import pdfplumber

        with pdfplumber.open(str(order)) as pdf:
            for number, page in enumerate(pdf.pages, start=1):
                if number > 1:
                    from docx.enum.text import WD_BREAK

                    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                document.add_paragraph(page.extract_text() or '')
        document.save(str(output))

        for mode in ('layout-editable', 'semantic-editable'):
            result = verify.check(order, output, mode=mode)
            assert result['status'] == 'defects', mode
            assert 'images' in checks(result), mode
            assert 'qr' in checks(result), mode

    def test_it_says_when_it_could_not_read_a_source_image(self, order, tmp_path, monkeypatch):
        output = page_image_docx(tmp_path / 'unread.docx', order)
        monkeypatch.setattr(probe, 'pdf_media', lambda path: [])
        detail = ' '.join(f['detail'] for f in verify.check(order, output)['findings'])
        assert 'could not be' in detail or 'has none' in detail

    def test_the_same_file_passes_as_visual_fidelity(self, order, tmp_path):
        output = page_image_docx(tmp_path / 'images2.docx', order)
        assert verify.check(order, output, mode='visual-fidelity')['status'] == 'clean'

    @needs_soffice
    def test_a_missing_page_is_a_defect(self, letter, order, tmp_path):
        output = tmp_path / 'one.docx'
        convert.convert(letter, output)
        assert 'pages' in checks(verify.check(order, output))

    @needs_soffice
    def test_a_lost_qr_is_a_defect(self, letter, tmp_path):
        output = tmp_path / 'letter.docx'
        convert.convert(letter, output, repair=False)
        stripped = _without_media(output, tmp_path / 'noqr.docx')
        assert 'qr' in checks(verify.check(letter, stripped))

    @needs_soffice
    def test_a_resized_page_is_a_defect(self, converted, order, tmp_path):
        output, _ = converted
        resized = tmp_path / 'resized.docx'
        _rewrite_page_size(output, resized)
        assert 'page_size' in checks(verify.check(order, resized))

    def test_the_floor_is_configurable_but_the_shortfall_is_not_hidden(self, order, tmp_path):
        output = tmp_path / 'thin.docx'
        document = Document()
        section = document.sections[0]
        section.page_width, section.page_height = Pt(WIDTH), Pt(HEIGHT)
        document.add_paragraph('Subject: Office order regarding the deployment.')
        document.save(str(output))
        assert 'text' in checks(verify.check(order, output))
        assert 'text' in checks(verify.check(order, output, min_text_ratio=0.01))


def _rewrite_page_size(source, target):
    with zipfile.ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    document = parts['word/document.xml'].decode('utf8')
    parts['word/document.xml'] = document.replace('w:w="11906"', 'w:w="12240"').encode('utf8')
    with zipfile.ZipFile(target, 'w') as output:
        for name, payload in parts.items():
            output.writestr(name, payload)
    return target


class TestCommandLine:
    @needs_soffice
    def test_a_clean_run_exits_zero(self, converted, order):
        output, _ = converted
        result = run(order, output)
        assert result.returncode == 0
        assert result.stdout.startswith('clean:')

    def test_defects_exit_two(self, order, tmp_path):
        output = page_image_docx(tmp_path / 'cli.docx', order)
        result = run(order, output)
        assert result.returncode == 2
        assert 'picture_book' in result.stdout

    def test_bad_input_exits_one(self, order, tmp_path):
        result = run(order, tmp_path / 'absent.docx')
        assert result.returncode == 1
        assert 'File not found' in result.stdout

    @needs_soffice
    def test_json_is_machine_readable(self, converted, order):
        output, _ = converted
        result = run(order, output, '--json')
        assert result.returncode == 0
        assert json.loads(result.stdout)['status'] == 'clean'

    def test_an_impossible_floor_is_rejected(self, order, tmp_path):
        output = page_image_docx(tmp_path / 'floor.docx', order)
        result = run(order, output, '--min-text', '1.5')
        assert result.returncode == 1
