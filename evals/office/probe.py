"""Inspect delivered office files and render them to PNG for review.

Usage: python probe.py --renders <dir> [--max-pages N] FILE...
Prints a JSON list with one report per file. Run through uv (see grade.ts) so the
document libraries never have to be installed globally.
"""

import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

TEXT_LIMIT = 20000
PARAGRAPH_LIMIT = 600
ERROR_VALUES = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
SOFFICE_CANDIDATES = [
    os.environ.get("SOFFICE", ""),
    shutil.which("soffice") or "",
    shutil.which("libreoffice") or "",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
]


def soffice():
    return next((path for path in SOFFICE_CANDIDATES if path and Path(path).exists()), None)


def base_report(path):
    return {
        "path": str(path),
        "name": path.name,
        "kind": "other",
        "bytes": path.stat().st_size,
        "ok": True,
        "text": "",
        "paragraphs": [],
        "fonts": [],
        "renders": [],
    }


def finish_text(report, paragraphs):
    cleaned = [p.strip() for p in paragraphs if p and p.strip()]
    report["paragraphs"] = cleaned[:PARAGRAPH_LIMIT]
    report["text"] = "\n".join(cleaned)[:TEXT_LIMIT]


def inherited_font(style):
    """The font a paragraph style sets, following base styles; None when nothing sets one."""
    while style is not None:
        font = getattr(style, "font", None)
        if font is not None and font.name:
            return font.name
        style = style.base_style
    return None


def inspect_docx(path, report):
    import docx

    document = docx.Document(str(path))
    paragraphs = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))
    fonts, outline = set(), []
    for paragraph in document.paragraphs:
        fonts.update(run.font.name for run in paragraph.runs if run.font.name)
        style_font = inherited_font(paragraph.style)
        if style_font:
            fonts.add(style_font)
        if paragraph.text.strip():
            style_name = paragraph.style.name if paragraph.style is not None else "(no style)"
            outline.append(f"{style_name}: {paragraph.text.strip()[:80]}")
    report["fonts"] = sorted(fonts)
    report["outline"] = outline[:200]
    report["images"] = len(document.inline_shapes)
    report["tables"] = len(document.tables)
    finish_text(report, paragraphs)


def inspect_pptx(path, report):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    deck = Presentation(str(path))
    paragraphs, titles, fonts, images = [], [], set(), 0
    for slide in deck.slides:
        title = slide.shapes.title.text_frame.text if slide.shapes.title is not None else ""
        titles.append(title.strip())
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                images += 1
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                paragraphs.append("".join(run.text for run in paragraph.runs))
                fonts.update(run.font.name for run in paragraph.runs if run.font.name)
    report["slides"] = len(deck.slides)
    report["slide_titles"] = titles
    report["images"] = images
    report["fonts"] = sorted(fonts)
    finish_text(report, paragraphs)


def inspect_xlsx(path, report):
    import openpyxl

    formulas_book = openpyxl.load_workbook(str(path), data_only=False)
    values_book = openpyxl.load_workbook(str(path), data_only=True)
    sheets, paragraphs, fonts = [], [], set()
    for sheet in formulas_book.worksheets:
        values = values_book[sheet.title]
        cells = formulas = errors = uncached = 0
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                cells += 1
                if cell.font is not None and cell.font.name:
                    fonts.add(cell.font.name)
                cached = values[cell.coordinate].value
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas += 1
                    if cached is None:
                        uncached += 1
                if isinstance(cached, str) and cached.strip() in ERROR_VALUES:
                    errors += 1
                paragraphs.append(str(cached if cached is not None else cell.value))
        sheets.append(
            {
                "name": sheet.title,
                "dims": sheet.dimensions,
                "cells": cells,
                "formulas": formulas,
                "errors": errors,
                "uncached": uncached,
                "charts": len(getattr(sheet, "_charts", [])),
            }
        )
    report["sheets"] = sheets
    report["fonts"] = sorted(fonts)
    finish_text(report, paragraphs)


def inspect_pdf(path, report):
    import fitz

    with fitz.open(str(path)) as document:
        report["pages"] = document.page_count
        fonts = set()
        paragraphs = []
        for page in document:
            paragraphs.extend(page.get_text().split("\n"))
            fonts.update(font[3] for font in page.get_fonts())
        report["fonts"] = sorted(fonts)
        finish_text(report, paragraphs)


def inspect_image(path, report):
    from PIL import Image, ImageStat

    with Image.open(str(path)) as image:
        report["width"], report["height"] = image.size
        grey = image.convert("L")
        report["blank"] = ImageStat.Stat(grey).stddev[0] < 3


def render_pdf(pdf, out_dir, stem, max_pages):
    import fitz

    paths = []
    with fitz.open(str(pdf)) as document:
        for index, page in enumerate(document):
            if index >= max_pages:
                break
            target = out_dir / f"{stem}-p{index + 1}.png"
            page.get_pixmap(dpi=80).save(str(target))
            paths.append(str(target))
    return paths


def render(path, report, out_dir, max_pages):
    stem = path.stem.replace(" ", "_")
    if report["kind"] == "pdf":
        report["renders"] = render_pdf(path, out_dir, stem, max_pages)
        return
    if report["kind"] == "image":
        report["renders"] = [str(path)]
        return
    binary = soffice()
    if binary is None:
        report["render_error"] = "LibreOffice not found: install it or set SOFFICE to render pages"
        return
    with tempfile.TemporaryDirectory() as tmp:
        profile = Path(tmp) / "profile"
        result = subprocess.run(
            [
                binary,
                f"-env:UserInstallation=file://{profile}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                tmp,
                str(path),
            ],
            capture_output=True,
            timeout=180,
        )
        pdf = Path(tmp) / f"{path.stem}.pdf"
        if result.returncode != 0 or not pdf.exists():
            report["render_error"] = f"LibreOffice could not convert the file: {result.stderr.decode()[:300]}"
            return
        report["renders"] = render_pdf(pdf, out_dir, stem, max_pages)


INSPECTORS = {
    ".docx": ("docx", inspect_docx),
    ".pptx": ("pptx", inspect_pptx),
    ".xlsx": ("xlsx", inspect_xlsx),
    ".xlsm": ("xlsx", inspect_xlsx),
    ".pdf": ("pdf", inspect_pdf),
}


def inspect(path, out_dir, max_pages):
    report = base_report(path)
    suffix = path.suffix.lower()
    kind, inspector = INSPECTORS.get(suffix, ("image", inspect_image) if suffix in IMAGE_EXTS else ("other", None))
    report["kind"] = kind
    if inspector is None:
        return report
    try:
        inspector(path, report)
    except Exception as error:  # noqa: BLE001 - an unreadable file is a result, not a crash
        report["ok"] = False
        report["error"] = f"{type(error).__name__}: {error}"
        return report
    try:
        render(path, report, out_dir, max_pages)
    except Exception as error:  # noqa: BLE001
        report["render_error"] = f"{type(error).__name__}: {error}"
    return report


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--renders", required=True)
    parser.add_argument("--max-pages", type=int, default=8)
    parser.add_argument("files", nargs="+")
    args = parser.parse_args()
    out_dir = Path(args.renders)
    out_dir.mkdir(parents=True, exist_ok=True)
    reports = [inspect(Path(file), out_dir, args.max_pages) for file in args.files]
    json.dump(reports, sys.stdout, ensure_ascii=False)


if __name__ == "__main__":
    main()
